from __future__ import annotations

import math
import os
import shutil
import sys
import textwrap
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
WORK = Path(sys.argv[3])
WORK.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F5F8"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
WHITE = "FFFFFF"
TEXT = "1F2937"


def font(size: int, bold: bool = False):
    path = Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size=size)


def fit_lines(draw: ImageDraw.ImageDraw, value: str, width: int, face, max_lines: int = 4):
    words = value.split()
    lines = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        if draw.textbbox((0, 0), test, font=face)[2] <= width or not current:
            current = test
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        while draw.textbbox((0, 0), lines[-1] + "…", font=face)[2] > width and lines[-1]:
            lines[-1] = lines[-1][:-1]
        lines[-1] += "…"
    return lines


def centered_text(draw, rect, title, subtitle=None, fill=TEXT, title_size=31, sub_size=23):
    x0, y0, x1, y1 = rect
    tf = font(title_size, True)
    sf = font(sub_size, False)
    lines = fit_lines(draw, title, x1 - x0 - 36, tf, 3)
    sub_lines = fit_lines(draw, subtitle, x1 - x0 - 36, sf, 3) if subtitle else []
    line_h = title_size + 8
    sub_h = sub_size + 7
    total = len(lines) * line_h + (12 if sub_lines else 0) + len(sub_lines) * sub_h
    y = y0 + (y1 - y0 - total) / 2
    for line in lines:
        box = draw.textbbox((0, 0), line, font=tf)
        draw.text(((x0 + x1 - (box[2] - box[0])) / 2, y), line, font=tf, fill=fill)
        y += line_h
    y += 6
    for line in sub_lines:
        box = draw.textbbox((0, 0), line, font=sf)
        draw.text(((x0 + x1 - (box[2] - box[0])) / 2, y), line, font=sf, fill=fill)
        y += sub_h


def box(draw, rect, title, subtitle=None, fill="#FFFFFF", outline="#2F75B5", title_fill="#17365D"):
    draw.rounded_rectangle(rect, radius=22, fill=fill, outline=outline, width=4)
    centered_text(draw, rect, title, subtitle, fill=title_fill)


def arrow(draw, start, end, color="#5B6573", width=6):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    head = 16
    left = (end[0] - head * math.cos(angle - 0.55), end[1] - head * math.sin(angle - 0.55))
    right = (end[0] - head * math.cos(angle + 0.55), end[1] - head * math.sin(angle + 0.55))
    draw.polygon([end, left, right], fill=color)


def canvas(title: str, subtitle: str):
    img = Image.new("RGB", (1600, 900), "#F7F9FC")
    draw = ImageDraw.Draw(img)
    draw.text((60, 42), title, font=font(42, True), fill="#17365D")
    draw.text((60, 101), subtitle, font=font(26), fill="#52606D")
    draw.line((60, 146, 1540, 146), fill="#A8BCD0", width=3)
    return img, draw


def make_figure_1(path: Path):
    img, d = canvas("Figure 1. Luồng Hybrid Router của AIFC Version 2", "Rule xử lý case rõ; Gemini chỉ được gọi khi bằng chứng chưa đủ")
    box(d, (55, 260, 260, 440), "Câu hỏi + lịch sử", "q, H, previous_subject", "#EAF2F8")
    box(d, (320, 260, 535, 440), "Rule Router", "keyword score, confidence, margin", "#E2F0D9", "#548235")
    box(d, (600, 210, 840, 350), "Đủ tin cậy?", "C ≥ 0.80; margin ≥ 0.25", "#FFF2CC", "#BF9000")
    box(d, (600, 470, 840, 650), "Semantic Router", "Gemini 2.5 Flash; JSON schema", "#FCE4D6", "#C55A11")
    box(d, (915, 260, 1135, 440), "Hybrid decision", "accept / override / clarify", "#DDEBF7")
    arrow(d, (260, 350), (320, 350))
    arrow(d, (535, 350), (600, 280))
    arrow(d, (720, 350), (720, 470))
    d.text((738, 375), "Không", font=font(22, True), fill="#C55A11")
    arrow(d, (840, 280), (915, 320))
    d.text((854, 238), "Có", font=font(22, True), fill="#548235")
    arrow(d, (840, 560), (915, 390))
    pool = (1195, 195, 1540, 680)
    d.rounded_rectangle(pool, radius=24, fill="#FFFFFF", outline="#5B9BD5", width=4)
    d.text((1288, 214), "MODEL MAP", font=font(27, True), fill="#17365D")
    models = [
        ("ENGLISH", "anhdai312/English19"),
        ("MATH", "anhdai312/Math19"),
        ("HISTORY", "anhdai312/History19"),
        ("GENERAL", "anhdai312/General19"),
    ]
    y = 275
    for label, model in models:
        fill = "#E2F0D9" if label != "GENERAL" else "#E4DFEC"
        box(d, (1230, y, 1505, y + 78), label, model, fill, "#7F8FA4")
        y += 92
    arrow(d, (1135, 350), (1195, 350))
    box(d, (480, 735, 1120, 840), "Response + structured log", "route, selected model, latency, token usage, raw response", "#EAF2F8")
    arrow(d, (1365, 680), (1080, 735))
    img.save(path)


def make_figure_2(path: Path):
    img, d = canvas("Figure 2. Thiết kế dữ liệu Router v2", "Calibration và locked test độc lập; ba môn cân bằng")
    box(d, (80, 205, 720, 420), "Calibration / development", "90 case tổng hợp · 30 English · 30 Math · 30 History", "#E2F0D9", "#548235")
    box(d, (880, 205, 1520, 420), "Locked hold-out test", "60 case tổng hợp · 20 English · 20 Math · 20 History", "#DDEBF7", "#2F75B5")
    arrow(d, (720, 312), (880, 312))
    d.text((727, 268), "không trùng câu hỏi", font=font(22, True), fill="#52606D")
    labels = [("Standard", 24, "#5B9BD5"), ("Typo", 9, "#ED7D31"), ("Abbreviation", 9, "#A5A5A5"), ("Follow-up", 9, "#70AD47"), ("Ambiguous", 9, "#FFC000")]
    d.text((82, 500), "Phân bố challenge trong locked test", font=font(30, True), fill="#17365D")
    x, y = 82, 565
    usable = 1435
    for label, count, color in labels:
        w = round(usable * count / 60)
        d.rectangle((x, y, x + w, y + 80), fill=color)
        if w > 170:
            centered_text(d, (x, y, x + w, y + 80), f"{label}: {count}", fill="#FFFFFF", title_size=23)
        x += w
    y2 = 690
    x = 82
    for label, count, color in labels:
        d.rectangle((x, y2 + 5, x + 26, y2 + 31), fill=color)
        d.text((x + 38, y2), f"{label} ({count})", font=font(22), fill="#374151")
        x += 275
    box(d, (240, 760, 1360, 850), "Gold labels", "gold_subject · gold_intent · gold_need_clarification · history · previous_subject", "#FFFFFF", "#7F8FA4")
    img.save(path)


def make_figure_3(path: Path):
    img, d = canvas("Figure 3. Kiến trúc triển khai", "Tách Router Benchmark khỏi suy luận SLM và Model Eval")
    box(d, (60, 260, 280, 430), "Web UI", "upload test, chọn mode, model map", "#EAF2F8")
    box(d, (340, 260, 590, 430), "Backend API", "/router/evaluate", "#DDEBF7")
    box(d, (655, 190, 925, 360), "Routing service", "Rule + Hybrid orchestrator", "#E2F0D9", "#548235")
    box(d, (655, 470, 925, 640), "Gemini Router", "OpenRouter provider; case khó", "#FCE4D6", "#C55A11")
    box(d, (1000, 190, 1250, 360), "Model Registry", "English19 · Math19 · History19 · General19", "#E4DFEC", "#8064A2")
    box(d, (1000, 470, 1250, 640), "GPU inference", "chỉ dùng ở End-to-End", "#FFF2CC", "#BF9000")
    box(d, (1320, 300, 1540, 530), "Artifacts", "predictions, metrics, replay, logs", "#FFFFFF", "#5B9BD5")
    arrow(d, (280, 345), (340, 345))
    arrow(d, (590, 345), (655, 275))
    arrow(d, (790, 360), (790, 470))
    arrow(d, (925, 275), (1000, 275))
    arrow(d, (925, 555), (1000, 555))
    arrow(d, (1250, 275), (1320, 360))
    arrow(d, (1250, 555), (1320, 470))
    d.text((85, 720), "RP4: /router/evaluate không sinh câu trả lời từ SLM.", font=font(27, True), fill="#17365D")
    d.text((85, 770), "RP5: Pooled / Oracle / Hybrid replay mới gọi GPU và Gemini Judge.", font=font(27, True), fill="#17365D")
    img.save(path)


def make_figure_4(path: Path):
    img, d = canvas("Figure 4. Giao thức đánh giá hai tầng", "Một locked test, hai câu hỏi đánh giá khác nhau")
    box(d, (50, 310, 255, 500), "Locked test", "60 routing cases", "#EAF2F8")
    box(d, (330, 180, 610, 360), "RP4 · Router", "RULE · LLM · HYBRID", "#DDEBF7")
    box(d, (330, 560, 610, 740), "RP5 · End-to-End", "POOLED · ORACLE · HYBRID", "#E2F0D9", "#548235")
    arrow(d, (255, 405), (330, 270))
    arrow(d, (255, 405), (330, 650))
    box(d, (690, 180, 1025, 360), "Router metrics", "accuracy · macro-F1 · intent · exact · latency · LLM call rate", "#FFFFFF")
    box(d, (690, 560, 1025, 740), "Generate + Judge", "raw response · A/B/C/D · latency · cost", "#FFFFFF", "#548235")
    arrow(d, (610, 270), (690, 270))
    arrow(d, (610, 650), (690, 650))
    box(d, (1110, 180, 1540, 360), "Bằng chứng RP4", "benchmark JSON + confusion matrix + error breakdown", "#F2F5F8", "#7F8FA4")
    box(d, (1110, 560, 1540, 740), "Bằng chứng RP5", "replay + route trace + Gemini/Human rubric", "#F2F5F8", "#7F8FA4")
    arrow(d, (1025, 270), (1110, 270))
    arrow(d, (1025, 650), (1110, 650))
    d.text((180, 815), "Không suy điểm Socratic/factuality từ kết quả Router.", font=font(29, True), fill="#C00000")
    img.save(path)


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_paragraph(paragraph, value: str, bold: bool | None = None):
    first = paragraph.runs[0] if paragraph.runs else None
    fmt = {
        "name": first.font.name if first else None,
        "size": first.font.size if first else None,
        "bold": first.bold if first else None,
        "italic": first.italic if first else None,
    }
    clear_paragraph(paragraph)
    run = paragraph.add_run(value)
    if fmt["name"]:
        run.font.name = fmt["name"]
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), fmt["name"])
    if fmt["size"]:
        run.font.size = fmt["size"]
    run.bold = fmt["bold"] if bold is None else bold
    run.italic = fmt["italic"]


def find_first(doc: Document, prefix: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"Paragraph not found: {prefix}")


def replace(doc: Document, prefix: str, value: str):
    set_paragraph(find_first(doc, prefix), value)


def remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def insert_after(paragraph, text: str, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    p.add_run(text)
    return p


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width_twips))
    tcW.set(qn("w:type"), "dxa")


def style_table(table, widths=None, font_size=8.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths) if widths else 9000
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    if widths:
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(width))
            grid.append(col)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if widths:
                set_cell_width(cell, widths[c_idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, NAVY if r_idx == 0 else (LIGHT_GRAY if r_idx % 2 == 0 else WHITE))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if c_idx > 0 and len(cell.text) < 22:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(255, 255, 255) if r_idx == 0 else RGBColor(31, 41, 55)
                    run.bold = r_idx == 0
        if r_idx == 0:
            trPr = row._tr.get_or_add_trPr()
            tblHeader = OxmlElement("w:tblHeader")
            tblHeader.set(qn("w:val"), "true")
            trPr.append(tblHeader)


def rewrite_table(table, data, widths=None, font_size=8.5):
    while len(table.rows) < len(data):
        table.add_row()
    while len(table.rows) > len(data):
        table._tbl.remove(table.rows[-1]._tr)
    for r, row in enumerate(data):
        while len(table.rows[r].cells) < len(row):
            raise ValueError("Cannot add table columns with this helper")
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = str(value)
    style_table(table, widths=widths, font_size=font_size)


def add_table_before(doc, ref, data, widths, font_size=8.5):
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    style_table(table, widths=widths, font_size=font_size)
    ref._p.addprevious(table._tbl)
    return table


def add_caption_before(ref, text):
    p = ref.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    return p


def patch_media(docx_path: Path, replacements: dict[str, Path]):
    temp = docx_path.with_suffix(".media-patched.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = replacements[item.filename].read_bytes() if item.filename in replacements else zin.read(item.filename)
            zout.writestr(item, data)
    os.replace(temp, docx_path)


def main():
    fig1 = WORK / "figure1_hybrid_pipeline.png"
    fig2 = WORK / "figure2_router_dataset.png"
    fig3 = WORK / "figure3_system_architecture.png"
    fig4 = WORK / "figure4_evaluation_protocol.png"
    make_figure_1(fig1)
    make_figure_2(fig2)
    make_figure_3(fig3)
    make_figure_4(fig4)

    doc = Document(SOURCE)

    # Abbreviations and change log.
    doc.tables[0].cell(7, 0).text = "H1–H5"
    doc.tables[0].cell(7, 1).text = "Hypothesis 1 to 5"
    doc.tables[1].cell(2, 0).text = "19/07/2026"
    doc.tables[1].cell(2, 1).text = "M"
    doc.tables[1].cell(2, 2).text = "Project team"
    doc.tables[1].cell(2, 3).text = "Aligned RP4 with Router v2 dataset, current implementation, preliminary 60-case benchmark, and reproducible evaluation limits."

    replace(doc, "Nghiên cứu sử dụng thiết kế thực nghiệm so sánh.",
            "Nghiên cứu sử dụng thiết kế thực nghiệm so sánh ở hai tầng. Trong RP4, Rule-only, LLM-only và Hybrid Router được chạy trên cùng locked routing test set để đo khả năng chọn đúng môn, ý định, nhu cầu làm rõ và chi phí định tuyến. Trong RP5, Pooled General, Oracle specialists và Hybrid end-to-end mới được so sánh về chất lượng câu trả lời, độ trễ sinh, token cost và điểm sư phạm.")
    replace(doc, "Khung giải pháp đề xuất được thiết kế",
            "Khung giải pháp đề xuất được module hóa. Truy vấn hiện tại q, lịch sử gần đây H và previous_subject được đưa vào Rule Router trước. Nếu confidence và decision margin vượt ngưỡng, hệ thống dùng ngay quyết định rule; nếu không, Gemini 2.5 Flash đóng vai trò Semantic Router. Hybrid Orchestrator hòa giải hai tín hiệu, sau đó ánh xạ ENGLISH, MATH, HISTORY hoặc GENERAL tới model active trong Registry.")
    replace(doc, "Kiến trúc cascade —",
            "Kiến trúc cascade — trong đó truy vấn được xử lý bởi thành phần chi phí thấp trước và chỉ chuyển sang thành phần phức tạp hơn khi cần — đã được chứng minh có hiệu quả trong việc cân bằng chi phí và chất lượng ở hệ thống đa mô hình [6]. RouteLLM cũng cho thấy dữ liệu ưu tiên có thể được dùng để học chính sách định tuyến [7]. Nghiên cứu hiện tại không huấn luyện RL Router; đóng góp ở cấp capstone là Hybrid cascade có ngưỡng, logging và protocol kiểm chứng tách biệt.")
    replace(doc, "Khung giải pháp này được lựa chọn",
            "Khung giải pháp này giải quyết khoảng trống của Version 1: một model chung phải xử lý mọi môn và nhãn môn học chưa được dùng để chọn mô hình lúc suy luận. Version 2 biến nhãn môn học thành quyết định có thể quan sát, benchmark và audit, với model map hiện tại gồm English19, Math19, History19 và General19.")
    replace(doc, "Tương tự cách tiếp cận cascaded pipeline",
            "Thiết kế module hóa cho phép kiểm tra riêng Rule Router, Semantic Router, Hybrid policy, từng SLM chuyên biệt và pipeline end-to-end. Nhờ đó lỗi chọn model không bị trộn với lỗi sinh câu trả lời, và Oracle condition có thể được dùng để ước lượng cận trên khi route luôn đúng.")
    replace(doc, "Nói cách khác, điểm nghiên cứu",
            "Điểm nghiên cứu không nằm ở số lượng model, mà ở việc đo xem Hybrid policy có giữ được độ chính xác gần LLM-only trong khi giảm số lần gọi LLM và độ trễ định tuyến hay không; sau đó RP5 mới kiểm tra liệu việc chọn đúng specialist có cải thiện chất lượng học tập/Socratic so với pooled model hay không.")

    replace(doc, "Tập dữ liệu chính bao gồm",
            "Dữ liệu Router v2 gồm hai phần tách biệt. router_train_v2.json là tập calibration/development tổng hợp gồm 90 case cân bằng (30 ENGLISH, 30 MATH, 30 HISTORY). router_test_v2.json là locked hold-out test gồm 60 case cân bằng (20 case mỗi môn). Tập test không được dùng để chọn prompt hoặc ngưỡng.")
    replace(doc, "Tập dữ liệu được sử dụng cho hai mục đích",
            "Mỗi test case có question, gold_subject, gold_intent và gold_need_clarification; case follow-up còn có history và previous_subject. Đây là dữ liệu kiểm tra Router, khác với dữ liệu messages dùng để fine-tune SLM sinh câu trả lời. Việc tách hai loại dữ liệu ngăn chỉ số định tuyến bị hiểu nhầm thành chất lượng phản hồi.")
    replace(doc, "Lý do chọn tập pilot 200 hội thoại",
            "Locked test được thiết kế có chủ đích để gây nhiễu ngôn ngữ: 24 case standard, 9 typo, 9 abbreviation, 9 follow-up và 9 ambiguous. Phân bố cân bằng theo môn giúp accuracy và Macro-F1 dễ diễn giải, còn các nhóm challenge kiểm tra độ bền trước câu viết tắt, sai chính tả, thiếu ngữ cảnh và câu hỏi nhiều lượt.")
    replace(doc, "Figure 2: Chèn ảnh", "Figure 2: Thiết kế calibration và locked test của Router v2.")
    replace(doc, "Figure 2: Cấu trúc dữ liệu", "Figure 4: Giao thức thu thập bằng chứng tách biệt cho RP4 và RP5.")
    image_paragraph_3 = doc.paragraphs[55]
    if not any(p.text.strip().startswith("Figure 3:") for p in doc.paragraphs):
        insert_after(image_paragraph_3, "Figure 3: Kiến trúc triển khai Router, Registry, GPU inference và artifact logger.", "normal")

    replace(doc, "• Thu thập các hội thoại", "• Soạn 90 case calibration tổng hợp và 60 case hold-out độc lập cho ba nhãn ENGLISH, MATH và HISTORY.")
    replace(doc, "• Gán hoặc xác minh nhãn", "• Gán thủ công gold_subject, gold_intent và gold_need_clarification; các case nhiều lượt được gắn history và previous_subject.")
    replace(doc, "• Tạo split manifest", "• Kiểm tra phân bố cân bằng theo môn, phân bố challenge và không trùng nguyên văn câu hỏi giữa calibration và test.")
    replace(doc, "• Giữ nguyên locked test set", "• Khóa router_test_v2.json trước benchmark; không sửa prompt, từ điển hoặc threshold dựa trên lỗi của tập này.")
    replace(doc, "• Lưu dữ liệu thô", "• Lưu test file, model map, threshold, mode, repetitions, prediction và metric trong artifact JSON để tái lập.")
    replace(doc, "• Chuẩn hóa schema hội thoại", "• Chuẩn hóa subject thành chữ hoa và kiểm tra các trường gold bắt buộc.")
    replace(doc, "• Loại bỏ artifact hệ thống", "• Chuẩn hóa khoảng trắng nhưng giữ nguyên typo/abbreviation có chủ đích.")
    replace(doc, "• Xác minh nhãn môn học", "• Rà soát thủ công nhãn subject, intent và clarification trước khi khóa test.")
    replace(doc, "• Tách dữ liệu huấn luyện theo môn", "• Giữ calibration data tách biệt với hold-out test; không dùng test để điều chỉnh keyword hoặc prompt.")
    replace(doc, "• Giữ nguyên locked test set để", "• Tính SHA-256 của danh sách case trong mỗi benchmark export để phát hiện thay đổi dữ liệu.")
    replace(doc, "Nghiên cứu sử dụng cách lấy mẫu phân tầng",
            "Nghiên cứu lấy mẫu phân tầng theo môn học: 20 ENGLISH, 20 MATH và 20 HISTORY trong locked test. Bên trong từng môn, các case được phân bổ vào năm nhóm challenge. Đây là purposive stratified sampling trên dữ liệu tổng hợp, không phải mẫu ngẫu nhiên đại diện cho toàn bộ học sinh; vì vậy kết luận chỉ có phạm vi pilot.")
    replace(doc, "Quá trình phân tích kết hợp log tự động",
            "Phân tích RP4 dùng log tự động từ Router Benchmark: predicted subject, intent, need_clarification, strategy, confidence, margin, LLM-called flag, latency, error type và confusion matrix. Rubric A1–D2 chỉ được áp dụng ở RP5 cho raw response do Pooled/Oracle/Hybrid sinh ra; không suy điểm Socratic hoặc factuality từ chỉ số Router.")
    replace(doc, "• Baseline V1: một model chung", "• Cấp Router (RP4): Rule-only và LLM-only là hai baseline; Hybrid là phương pháp đề xuất.")
    replace(doc, "• Proposed V2: Hybrid Router", "• Cấp End-to-End (RP5): Pooled General (B), Oracle specialists (C) và Hybrid Router + specialists (D).")
    replace(doc, "Kết quả baseline pilot cho thấy",
            "Baseline V1 ở bảng dưới là bằng chứng lịch sử về chất lượng một model chung trên pilot cũ, không phải baseline trực tiếp của benchmark Router English–Math–History. Kết luận Router trong RP4 chỉ dựa trên Rule, LLM và Hybrid chạy trên cùng router_test_v2.json; so sánh chất lượng thế hệ V1/V2 được chuyển sang RP5.")

    replace(doc, "Hệ thống Version 2 được đề xuất",
            "Version 2 là Hybrid Subject Router cho hội thoại giáo dục. Đầu vào gồm q, tối đa các lượt lịch sử gần nhất H và previous_subject khi có. Đầu ra gồm subject, secondary_subjects, intent, confidence, margin, needClarification, strategy, selectedModel, llmCalled và latencyMs. Quyết định route được benchmark độc lập trước khi gọi SLM sinh phản hồi.")
    replace(doc, "4.2 Thuật toán 2: Bộ định tuyến ngữ nghĩa dựa trên Qwen",
            "4.2 Thuật toán 2: Bộ định tuyến ngữ nghĩa dựa trên Gemini")
    replace(doc, "Luồng xử lý end-to-end là:",
            "Luồng xử lý là: q + H + previous_subject → Rule Router → confidence/margin gate → Gemini-based Semantic Router cho case khó → Hybrid reconciliation → model map → English19/Math19/History19 hoặc General19 → response và structured log. Trong Router Benchmark, luồng dừng ở quyết định route; GPU generation chỉ chạy ở End-to-End RP5.")
    replace(doc, "Đóng góp nghiên cứu.",
            "Đóng góp nghiên cứu là chính sách cascade tường minh và có thể audit: rule xử lý case có bằng chứng mạnh; LLM xử lý case thiếu từ khóa, typo, mơ hồ hoặc follow-up; xung đột không đủ mạnh dẫn đến clarification thay vì ép route sang specialist sai.")
    replace(doc, "Prompt của mô hình quy định rõ",
            "Prompt Semantic Router yêu cầu JSON-only theo schema cố định và giới hạn nhãn trong model map. Toán chỉ là công cụ trong bài thuộc môn khác không được tự động gán MATH. Với model map hiện tại, các nhãn mục tiêu là ENGLISH, MATH, HISTORY, GENERAL và UNKNOWN.")
    replace(doc, "Theo đó:", "Theo model map của campaign hiện tại:")
    replace(doc, "Route Math →", "Route ENGLISH → anhdai312/English19.")
    replace(doc, "Route Physics →", "Route MATH → anhdai312/Math19.")
    replace(doc, "Route Chemistry →", "Route HISTORY → anhdai312/History19.")
    replace(doc, "Route General hoặc truy vấn", "Route GENERAL hoặc nhãn không có specialist → anhdai312/General19; UNKNOWN/thiếu thông tin → AskClarification theo chính sách.")
    replace(doc, "Mỗi SLM chuyên biệt được fine-tune",
            "Bốn model registry active của campaign là General19, English19, Math19 và History19. Cấu hình training được ghi nhận gồm 3 epochs, batch size 1, gradient accumulation 4, learning rate 5×10⁻⁵, block size/max length 1024, warmup 5 steps, weight decay 0.01, optimizer adamw_8bit, linear scheduler, seed 3407, LoRA rank 16 và alpha 32. Đây là cấu hình thực nghiệm cần được báo cáo; ảnh hưởng của hyperparameter phải được đánh giá bằng validation loss và pilot generation, không suy luận từ train loss duy nhất.")
    replace(doc, "Phần triển khai sử dụng giao diện benchmark",
            "Giao diện Router Benchmark nhận wrapper JSON có trường cases, cho phép chọn RULE, LLM, HYBRID, repetitions và ablation. Backend chạy tuần tự từng mode/case, trả prediction, per-class metrics, confusion matrix, error breakdown và evaluation manifest. Full ablation trên 60 case làm tăng mạnh số lượt gọi Gemini; vì vậy campaign chính dùng repetitions = 1, còn stability/ablation chỉ nên chạy trên subset đã định trước hoặc bằng job có progress.")
    replace(doc, "Cấu hình triển khai được cố định",
            "Campaign sơ bộ được cố định với router_test_v2.json, model map English19/Math19/History19/General19, threshold mặc định θ_rule=0.80, μ_rule=0.25, θ_llm=0.70 và Semantic Router google/gemini-2.5-flash. Các threshold này là cấu hình triển khai hiện tại, chưa phải giá trị tối ưu được chứng minh; calibration/grid search phải dùng router_train_v2.json và khóa trước lần benchmark cuối.")

    replace(doc, "Chiến lược đánh giá nhằm trả lời",
            "Chiến lược đánh giá tách hai câu hỏi: Router có chọn đúng model với chi phí hợp lý hay không (RP4), và model được chọn có tạo phản hồi đúng/Socratic hơn pooled model hay không (RP5). Cả ba Router mode dùng cùng 60 locked cases; mọi kết luận ở §5.5 chỉ thuộc cấp Router.")
    replace(doc, "Chiến lược đánh giá được chia thành hai tầng.",
            "Tầng component-level đo Rule, LLM và Hybrid bằng subject accuracy, Macro-F1, intent accuracy, exact match, wrong-route rate, latency và LLM call rate. Tầng end-to-end giữ cùng replay nhưng bổ sung generation, Pooled/Oracle/Hybrid condition và Gemini/Human rubric; tầng này chưa được chạy trong số liệu RP4 hiện tại.")
    replace(doc, "Thiết kế ban đầu so sánh Version 1",
            "Thiết kế end-to-end gồm bốn nhóm có vai trò khác nhau. Nhóm A là baseline V1 lịch sử để cung cấp bối cảnh; nhóm B là General19/pooled control trong pipeline hiện tại; nhóm C dùng gold_subject để chọn specialist (Oracle); nhóm D dùng Hybrid Router thực tế. C và D phải dùng cùng specialist và decoding setting để routing penalty chỉ phản ánh sai số route.")
    replace(doc, "Để kiểm soát biến nhiễu này",
            "Sự khác nhau về số mẫu hoặc chất lượng dữ liệu fine-tune giữa General19 và các specialist là biến gây nhiễu phải được ghi rõ. Vì vậy kết luận chính RP5 cần báo cáo dataset manifest của từng model, train/validation split, base model, seed và decoding config; nếu chưa cân bằng được dữ liệu thì kết quả chỉ được gọi là pilot engineering evidence.")
    replace(doc, "Với thiết kế này, phép so sánh B với C/D",
            "B so với C đo lợi ích tối đa của chuyên biệt hóa khi route luôn đúng; C so với D đo routing penalty; B so với D là hiệu quả thực tế của hệ thống đề xuất. A không được trộn vào kiểm định chính nếu pipeline, dữ liệu và rubric khác với B/C/D.")

    # Remove the duplicated overview under the first 5.2 heading.
    first_52 = find_first(doc, "5.2 Chỉ số đánh giá và công thức")
    current = first_52._p.getnext()
    while current is not None:
        from docx.text.paragraph import Paragraph
        p = Paragraph(current, first_52._parent)
        if p.text.strip().startswith("5.2.1."):
            break
        nxt = current.getnext()
        if current.tag == qn("w:p"):
            current.getparent().remove(current)
        current = nxt
    duplicates = [p for p in doc.paragraphs if p.text.strip() == "5.2 Chỉ số đánh giá và công thức"]
    for p in duplicates[1:]:
        remove_paragraph(p)

    replace(doc, "Để đánh giá cân bằng giữa các môn học",
            "Để đánh giá cân bằng giữa các nhãn xuất hiện trong gold hoặc prediction, nghiên cứu sử dụng Macro-F1:")
    replace(doc, "S là tập tất cả các môn học.",
            "S_eval là hợp của nhãn gold và nhãn được dự đoán trong campaign. Vì vậy một nhãn ngoài phạm vi như UNKNOWN/CHEMISTRY nếu bị dự đoán sai vẫn nhận F1=0 và làm giảm Macro-F1.")
    replace(doc, "Nᶜˡᵃʳ là số lượng truy vấn",
            "Nˡᵃᵇ là số case có gold_need_clarification (trong locked test hiện tại là toàn bộ 60 case, gồm cả true và false).")
    replace(doc, "Ngoài các chỉ số trên, nghiên cứu còn báo cáo",
            "Exact Match chỉ đúng khi subject, intent và need_clarification đồng thời khớp gold label. Wrong-route rate = 1 − Primary Subject Accuracy. Nghiên cứu còn báo cáo confusion matrix và error breakdown của Rule, LLM và Hybrid trên cùng locked test set.")
    replace(doc, "Độ ổn định của hệ thống được xác định",
            "Độ ổn định Router được xác định từ decision signature gồm subject, intent, clarification và strategy qua nhiều lần lặp cùng cấu hình:")

    # Add current Model Eval formula and weight-selection protocol before 5.2.3.
    ref_523 = find_first(doc, "5.2.3. Chỉ số so sánh")
    p = ref_523.insert_paragraph_before("Model Eval trong RP5 sử dụng rubric A1–D2 trên thang 0–5. Các nhóm điểm được tính theo: A=0.5A1+0.3A2+0.2A3; B=0.6B1+0.4B2; C=0.4C1+0.4C2+0.2C3; D=0.5D1+0.5D2. Nếu A1≤1 thì A bị giới hạn không quá 1.0 để việc đưa đáp án trực tiếp không bị các tiêu chí khác bù trừ.")
    p.style = "normal"
    p = ref_523.insert_paragraph_before("Điểm tổng hợp đang được triển khai là Overall=0.40A+0.25B+0.25C+0.10D. A có trọng số cao nhất vì bảo toàn phương pháp Socratic là mục tiêu sư phạm cốt lõi; B và C cùng 0.25 để độ đúng kiến thức và chất lượng hướng dẫn vẫn là điều kiện bắt buộc; D nhận 0.10 vì hallucination đã được phản ánh một phần ở B/C, còn tốc độ là thuộc tính hệ thống thứ cấp. D2 được ánh xạ theo latency trung bình: ≤2 s: 5; ≤4 s: 4; ≤7 s: 3; ≤12 s: 2; >12 s: 1.")
    p.style = "normal"
    p = ref_523.insert_paragraph_before("Các trọng số trên là giả thuyết đo lường, không phải hằng số tùy ý. Trước kết luận RP5, nhóm phải: (i) khóa rubric trước test; (ii) kiểm tra tương quan với điểm người chấm trên một validation subset; (iii) chạy sensitivity analysis với ít nhất ba cấu hình: Socratic-priority (0.40/0.25/0.25/0.10), accuracy-priority (0.30/0.35/0.25/0.10) và balanced (0.30/0.30/0.30/0.10); (iv) chỉ kết luận mô hình thắng nếu thứ hạng không đảo chiều hoặc giải thích rõ độ nhạy.")
    p.style = "normal"

    replace(doc, "Các công thức trên chỉ mô tả",
            "Các công thức trên được áp dụng cho benchmark Router sơ bộ ở §5.5. Điểm chất lượng, Socratic, factuality và generation latency chưa xuất hiện trong RP4; chúng chỉ được tính sau khi có replay Pooled/Oracle/Hybrid ở RP5.")
    replace(doc, "H5: Các ngưỡng độ tin cậy",
            "H5: Threshold phải được chọn trên calibration/validation data, không dùng locked test. Campaign sơ bộ hiện dùng giá trị mặc định θ_rule=0.80, μ_rule=0.25 và θ_llm=0.70; báo cáo cuối phải bổ sung grid search trên router_train_v2.json hoặc ghi đây là giới hạn nếu chưa thực hiện.")
    replace(doc, "Bước 1: Cố định test file",
            "Bước 1: Cố định router_test_v2.json, subject-model map, threshold, Semantic Router model, repetitions và thời điểm chạy; lưu evaluation manifest.")
    replace(doc, "Bước 2: chạy component-level",
            "Bước 2: chạy RULE, LLM và HYBRID trên toàn bộ 60 locked cases; xuất prediction, confidence, latency, per-class metrics, confusion matrix và error breakdown.")
    replace(doc, "Bước 7: diễn giải",
            "Bước 7: diễn giải H1–H3 chỉ từ Router metrics ở RP4; H4 và chất lượng sư phạm chỉ được kết luận sau End-to-End/Model Eval ở RP5.")

    # Update limitations to the evidence actually available.
    replace(doc, "Quy mô dữ liệu:", "Quy mô và nguồn dữ liệu: locked test chỉ có 60 case tổng hợp. Kết quả là pilot evidence, chưa đại diện cho hội thoại học sinh thật.")
    replace(doc, "Phạm vi môn học:", "Phạm vi nhãn: test hiện chỉ có ENGLISH, MATH và HISTORY; chưa đo GENERAL, code-switching ngoài thiết kế hoặc các môn khác.")
    replace(doc, "Sai số router:", "Số lần lặp: campaign hiện dùng repetitions=1, vì vậy stability chưa được đo và không được phép ghi là 100%.")
    replace(doc, "Biến động phần cứng:", "Token/cost: provider không trả token usage trong bảng hiện tại, nên chưa thể khẳng định mức tiết kiệm chi phí dù LLM call rate giảm.")
    replace(doc, "Thiên lệch đánh giá:", "Độ trễ và phụ thuộc provider: latency Gemini/OpenRouter có thể thay đổi theo mạng và tải dịch vụ; cần lặp trên subset cố định và báo cáo trung bình/độ lệch chuẩn.")
    replace(doc, "Rủi ro đánh đổi:", "Giới hạn suy luận: Router accuracy không chứng minh specialist trả lời đúng hoặc Socratic. Chất lượng model hiện phải được kiểm tra riêng bằng Pooled–Oracle–Hybrid replay và Model Eval trong RP5.")

    # Rewrite the most important tables to match the implemented study.
    rewrite_table(doc.tables[2], [
        ["Thành phần", "Chức năng", "Vai trò trong nghiên cứu"],
        ["Conversation context", "Chuẩn hóa q, history gần đây và previous_subject.", "Kiểm tra follow-up nhiều lượt."],
        ["Rule Router", "Keyword score có trọng số; tính confidence và margin.", "Baseline nhanh, giải thích được."],
        ["Gemini Semantic Router", "Trả JSON subject/intent/confidence/clarification.", "Fallback cho case khó."],
        ["Hybrid Orchestrator", "Confidence gate, hòa giải xung đột, clarification.", "Phương pháp đề xuất."],
        ["Model Registry", "ENGLISH→English19; MATH→Math19; HISTORY→History19; GENERAL→General19.", "Ánh xạ route sang model active."],
        ["Artifact logger", "Lưu manifest, prediction, metric, latency và replay.", "Bằng chứng RP4/RP5."],
    ], [1800, 3600, 3600])
    rewrite_table(doc.tables[3], [
        ["Tiêu chí dataset", "Lý do"],
        ["Cân bằng 3 môn", "20 test case cho mỗi ENGLISH, MATH và HISTORY; tránh majority-class bias."],
        ["Gold labels đầy đủ", "Đánh giá subject, intent, clarification và exact match."],
        ["Challenge có chủ đích", "Kiểm tra standard, typo, abbreviation, follow-up và ambiguous."],
        ["Calibration/test tách biệt", "Không chọn prompt hoặc threshold trên locked test."],
        ["Artifact JSON", "Tái lập input, configuration và output benchmark."],
    ], [3000, 6000])
    rewrite_table(doc.tables[4], [
        ["Trường dữ liệu", "Mô tả", "Mục đích"],
        ["id", "Định danh case duy nhất", "Truy vết lỗi và artifact"],
        ["question", "Câu hỏi hiện tại", "Đầu vào Router"],
        ["history / previous_subject", "Ngữ cảnh cho case follow-up", "Đo đóng góp của context"],
        ["gold_subject", "ENGLISH, MATH hoặc HISTORY", "Subject accuracy/F1"],
        ["gold_intent", "Một trong sáu intent có nhãn", "Intent accuracy"],
        ["gold_need_clarification", "Nhãn boolean", "Clarification/exact match"],
        ["challenge_type", "standard/typo/abbreviation/follow_up/ambiguous", "Phân tích robustness"],
    ], [2200, 3400, 3400])
    rewrite_table(doc.tables[5], [
        ["Tập dữ liệu", "Quy mô", "Mục đích / ràng buộc"],
        ["router_train_v2.json", "90 (30/môn)", "Calibration keyword/prompt/threshold; cần rà soát nhãn."],
        ["router_test_v2.json", "60 (20/môn)", "Locked hold-out; không dùng để tuning."],
        ["SLM training datasets", "Tách theo model", "Fine-tune General19/English19/Math19/History19; không dùng để tính Router accuracy."],
    ], [2700, 1800, 4500])
    rewrite_table(doc.tables[9], [
        ["Route cuối", "Mô hình active", "Chính sách"],
        ["ENGLISH", "anhdai312/English19", "Rule hoặc Semantic/Hybrid được chấp nhận."],
        ["MATH", "anhdai312/Math19", "Rule hoặc Semantic/Hybrid được chấp nhận."],
        ["HISTORY", "anhdai312/History19", "Rule hoặc Semantic/Hybrid được chấp nhận."],
        ["GENERAL", "anhdai312/General19", "Greeting, ngoài phạm vi hoặc specialist chưa có."],
        ["UNKNOWN", "AskClarification", "Không ép route khi bằng chứng chưa đủ."],
    ], [1800, 3200, 4000])
    rewrite_table(doc.tables[10], [
        ["Model", "HF repo", "Vai trò"],
        ["General19 (B)", "anhdai312/General19", "Pooled/fallback control"],
        ["English19", "anhdai312/English19", "Specialist ENGLISH"],
        ["Math19", "anhdai312/Math19", "Specialist MATH"],
        ["History19", "anhdai312/History19", "Specialist HISTORY"],
    ], [2200, 3300, 3500])
    rewrite_table(doc.tables[12], [
        ["Cấu hình", "Giá trị hiện tại", "Trạng thái phương pháp"],
        ["Semantic Router", "google/gemini-2.5-flash", "Cố định cho campaign"],
        ["θ_rule", "0.80", "Default; cần calibration"],
        ["μ_rule", "0.25", "Default; cần calibration"],
        ["θ_llm", "0.70", "Default; cần calibration"],
        ["Primary repetitions", "1", "Không suy stability"],
        ["Ablation", "Không dùng trong bảng chính", "Chạy subset định trước để kiểm soát token/time"],
    ], [2600, 2600, 3800])
    rewrite_table(doc.tables[13], [
        ["Tầng", "Điều kiện", "Chỉ số chính", "Mục đích"],
        ["RP4 Router", "RULE", "Accuracy, Macro-F1, intent, exact, latency", "Baseline rẻ/giải thích được"],
        ["RP4 Router", "LLM", "Cùng metric + 100% LLM call", "Cận trên semantic theo chi phí"],
        ["RP4 Router", "HYBRID", "Cùng metric + call rate", "Đo trade-off accuracy–latency"],
        ["RP5 End-to-End", "POOLED / ORACLE / HYBRID", "A/B/C/D, latency, token, stability", "Đo chất lượng câu trả lời"],
    ], [1700, 2100, 3000, 2200])
    rewrite_table(doc.tables[14], [
        ["Nhóm", "Kiến trúc", "Vai trò so sánh", "Giới hạn"],
        ["A", "Version 1 lịch sử", "Bối cảnh baseline cũ", "Không cùng pipeline/dataset hiện tại"],
        ["B", "General19 / pooled", "Control sinh câu trả lời", "Cần khóa dataset manifest"],
        ["C", "Oracle specialists", "Gold subject chọn specialist; cận trên", "Không phải hệ thống deploy thực tế"],
        ["D", "Hybrid + specialists", "Hệ thống đề xuất", "Chịu routing penalty"],
    ], [900, 2700, 3100, 2300])
    rewrite_table(doc.tables[15], [
        ["Giai đoạn", "Quy trình", "Đầu ra"],
        ["1", "Khóa test, model map, threshold, provider và repetitions.", "Evaluation manifest"],
        ["2", "Chạy RULE, LLM, HYBRID trên 60 case.", "RP4 Router metrics + errors"],
        ["3", "Calibration/ablation trên development hoặc subset định trước.", "Threshold/sensitivity evidence"],
        ["4", "Chạy POOLED và ORACLE replay.", "RP5 control + upper bound"],
        ["5", "Chạy HYBRID End-to-End.", "Replay D + route trace"],
        ["6", "Gemini/Human Judge và phân tích thống kê.", "RP5 quality evidence"],
    ], [1400, 4800, 2800])

    # Add preliminary benchmark section before limitations.
    ref_lim = find_first(doc, "6. Giới hạn của phương pháp nghiên cứu")
    h = ref_lim.insert_paragraph_before("5.5 Kết quả kiểm chứng Router sơ bộ trên 60 locked cases")
    h.style = "Heading 3"
    h.paragraph_format.page_break_before = True
    p = ref_lim.insert_paragraph_before("Bảng dưới ghi lại kết quả hiển thị từ campaign ngày 19/07/2026 trên router_test_v2.json. Ba mode dùng cùng 60 case, threshold mặc định và repetitions=1. Phần trăm được làm tròn theo giao diện; artifact tổng hợp được lưu tại docs/experiments/router_benchmark_summary_2026-07-19.json.")
    p.style = "normal"
    add_caption_before(ref_lim, "Bảng 16. Hiệu quả dự đoán của ba Router mode")
    add_table_before(doc, ref_lim, [
        ["Mode", "Subject Acc.", "Macro-F1", "Intent Acc.", "Exact Match", "Wrong-route"],
        ["RULE", "22%", "20%", "18%", "7%", "78%"],
        ["LLM", "77%", "65%", "52%", "43%", "23%"],
        ["HYBRID", "78%", "52%", "52%", "45%", "22%"],
    ], [1300, 1550, 1450, 1550, 1550, 1600], 8.5)
    add_caption_before(ref_lim, "Bảng 17. Chi phí định tuyến quan sát được")
    add_table_before(doc, ref_lim, [
        ["Mode", "Router latency", "LLM call rate", "LLM calls", "Stability", "Router tokens"],
        ["RULE", "2 ms", "0%", "0/60", "Chưa đo", "Không có"],
        ["LLM", "1,142 ms", "100%", "60/60", "Chưa đo", "Không có"],
        ["HYBRID", "823 ms", "70%", "42/60", "Chưa đo", "Không có"],
    ], [1300, 1700, 1700, 1300, 1500, 1500], 8.5)
    add_caption_before(ref_lim, "Bảng 18. Recall và F1 theo môn (support = 20/môn)")
    add_table_before(doc, ref_lim, [
        ["Mode", "English R/F1", "History R/F1", "Math R/F1", "Error breakdown chính"],
        ["RULE", "40% / 57%", "20% / 33%", "5% / 10%", "41 abstention; 6 misroute"],
        ["LLM", "90% / 95%", "60% / 75%", "80% / 89%", "14 abstention"],
        ["HYBRID", "95% / 97%", "60% / 75%", "80% / 89%", "7 abstention; 6 rule misroute"],
    ], [1100, 1700, 1700, 1700, 2800], 8.2)
    p = ref_lim.insert_paragraph_before("Kết quả cho thấy Hybrid tăng 56 điểm phần trăm subject accuracy so với Rule (78% so với 22%) và giảm wrong-route từ 78% xuống 22%. So với LLM-only, Hybrid tăng 1 điểm phần trăm accuracy, đồng thời giảm 18 lượt gọi LLM (30%) và giảm latency trung bình khoảng 27.9% (từ 1,142 ms xuống 823 ms). Vì vậy H1 được hỗ trợ trong phạm vi pilot; H2 được hỗ trợ về LLM call rate và latency, nhưng chưa được hỗ trợ về token cost do provider không trả usage.")
    p.style = "normal"
    p = ref_lim.insert_paragraph_before("Macro-F1 của Hybrid (52%) thấp hơn LLM-only (65%) dù subject accuracy cao hơn. Nguyên nhân là implementation tính Macro-F1 trên hợp của gold và predicted labels; các dự đoán CHEMISTRY/UNKNOWN ngoài ba gold classes nhận F1=0 và bị tính vào trung bình. Đây không phải mâu thuẫn số liệu mà là hình phạt cho out-of-scope prediction. Report cuối nên trình bày đồng thời Macro-F1 trên fixed gold label set và out-of-scope rate để người đọc diễn giải rõ hơn.")
    p.style = "normal"
    p = ref_lim.insert_paragraph_before("Rule Router có precision 100% trên ba lớp khi đưa ra đúng nhãn nhưng recall rất thấp, cho thấy hệ thống quá thường xuyên abstain/clarify khi thiếu từ khóa. Hybrid giảm incorrect clarification/abstention từ 41 xuống 7, nhưng vẫn còn 6 rule misroute. H3 chưa thể kết luận chắc chắn vì bảng hiện tại chưa xuất metric tách theo challenge_type; cần phân tích standard/typo/abbreviation/follow-up/ambiguous từ per-case export.")
    p.style = "normal"
    p = ref_lim.insert_paragraph_before("Kết luận RP4: Hybrid là lựa chọn triển khai hợp lý hơn Rule-only và tiết kiệm hơn LLM-only ở cấp Router. Tuy nhiên, kết quả này chưa chứng minh English19, Math19 hoặc History19 trả lời tốt hơn General19. Kết luận đó chỉ hợp lệ sau khi hoàn thành Pooled–Oracle–Hybrid End-to-End replay và Model Eval trong RP5.")
    p.style = "normal"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    patch_media(OUTPUT, {
        "word/media/image27.png": fig1,
        "word/media/image15.png": fig2,
        "word/media/image30.png": fig3,
        "word/media/image16.png": fig4,
    })
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
