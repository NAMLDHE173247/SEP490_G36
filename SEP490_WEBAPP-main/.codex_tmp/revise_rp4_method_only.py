from __future__ import annotations

import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])

NAVY = "17365D"
LIGHT_GRAY = "F2F5F8"
WHITE = "FFFFFF"


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_paragraph(paragraph, value: str, bold: bool | None = None):
    first = paragraph.runs[0] if paragraph.runs else None
    name = first.font.name if first else None
    size = first.font.size if first else None
    old_bold = first.bold if first else None
    clear_paragraph(paragraph)
    run = paragraph.add_run(value)
    if name:
        run.font.name = name
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = size
    run.bold = old_bold if bold is None else bold


def find_first(doc, prefix: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"Not found: {prefix}")


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width))
    tcW.set(qn("w:type"), "dxa")


def style_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, NAVY if row_idx == 0 else (LIGHT_GRAY if row_idx % 2 == 0 else WHITE))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if col_idx in (0, 3):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                    run.font.size = Pt(8.2)
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_idx == 0 else RGBColor(31, 41, 55)
                    run.bold = row_idx == 0 or (col_idx == 0 and row_idx > 0)
        if row_idx == 0:
            trPr = row._tr.get_or_add_trPr()
            tblHeader = OxmlElement("w:tblHeader")
            tblHeader.set(qn("w:val"), "true")
            trPr.append(tblHeader)


def add_table_before(ref, rows, widths):
    doc = ref.part.document
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    style_table(table, widths)
    ref._p.addprevious(table._tbl)
    return table


def add_heading_before(ref, text, style="Heading 4"):
    p = ref.insert_paragraph_before(text, style=style)
    return p


def add_text_before(ref, text):
    p = ref.insert_paragraph_before(text, style="normal")
    return p


def remove_result_section(doc):
    body = doc._body._element
    children = list(body)
    start = None
    end = None
    for index, child in enumerate(children):
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc._body)
            text = p.text.strip()
            if text.startswith("5.5 Kết quả kiểm chứng Router sơ bộ"):
                start = index
            if start is not None and text.startswith("6. Giới hạn của phương pháp nghiên cứu"):
                end = index
                break
    if start is None or end is None:
        raise ValueError("Could not locate RP4 result section")
    for child in children[start:end]:
        body.remove(child)


def remove_reference(doc, prefix):
    for p in list(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            p._element.getparent().remove(p._element)
            return True
    return False


def main():
    doc = Document(SOURCE)

    # Results belong in RP5. Keep only methodology in RP4.
    remove_result_section(doc)

    # Insert rubric definition table before the comparison metrics subsection.
    ref_523 = find_first(doc, "5.2.3. Chỉ số so sánh")
    add_heading_before(ref_523, "5.2.2.1. Định nghĩa Rubric A1–D2")
    add_text_before(ref_523, "Rubric được định nghĩa trước khi chạy Model Eval. Mỗi tiêu chí được chấm trên thang 0–5; điểm 0 thể hiện hoàn toàn không đạt, còn điểm 5 thể hiện đạt đầy đủ và nhất quán. D1 và D2 được chuẩn hóa theo hướng điểm cao hơn là tốt hơn.")
    rubric_rows = [
        ["Mã", "Tên tiêu chí", "Định nghĩa vận hành", "Thang điểm"],
        ["A1", "Answer Withholding", "Không cung cấp đáp án cuối trực tiếp khi học sinh chưa tự giải; chỉ xác nhận/gợi mở bước kế tiếp.", "0–5"],
        ["A2", "Scaffolding Quality", "Mức độ gợi mở từng bước, chia nhỏ nhiệm vụ và dẫn học sinh tự suy luận.", "0–5"],
        ["A3", "Adaptive Response", "Điều chỉnh phản hồi theo đúng/sai, lịch sử hội thoại, trình độ và trạng thái hiện tại của học sinh.", "0–5"],
        ["B1", "Factual Accuracy", "Công thức, phép tính, sự kiện và lập luận đúng; không có thông tin bịa đặt hoặc kết luận sai.", "0–5"],
        ["B2", "Grade-Level Appropriateness", "Ngôn ngữ, thuật ngữ, độ khó và cách giải thích phù hợp với cấp độ người học.", "0–5"],
        ["C1", "Robustness", "Duy trì chất lượng khi gặp typo, abbreviation, câu mơ hồ, follow-up hoặc chuyển ngữ cảnh.", "0–5"],
        ["C2", "Coherence", "Phản hồi nhất quán với câu hỏi và lịch sử; không tự mâu thuẫn, không lặp vô ích.", "0–5"],
        ["C3", "Tone", "Giọng điệu tôn trọng, khuyến khích, tự nhiên và phù hợp với vai trò gia sư.", "0–5"],
        ["D1", "Hallucination Control", "Không bịa dữ kiện, nguồn, công thức hoặc năng lực hệ thống; nếu thiếu thông tin thì nói rõ giới hạn.", "0–5"],
        ["D2", "Response Speed", "Điểm tốc độ theo latency trung bình của phản hồi: ≤2s=5; ≤4s=4; ≤7s=3; ≤12s=2; >12s=1.", "1–5"],
    ]
    add_table_before(ref_523, rubric_rows, [650, 2200, 5350, 800])
    add_text_before(ref_523, "Căn cứ chọn trọng số nhóm A–D. Tỷ lệ 0.5/0.3/0.2 cho A1/A2/A3 hiện được chọn dựa trên nhận định định tính của nhóm rằng Answer Withholding là điều kiện nhị phân quan trọng nhất trong việc bảo toàn phương pháp Socratic: nếu phản hồi đưa đáp án trực tiếp trước khi học sinh suy luận, giá trị gợi mở của lượt đó bị suy giảm mạnh. Scaffolding và Adaptation là các yếu tố bổ trợ. Đây là quyết định thiết kế ban đầu, chưa được kiểm định bằng dữ liệu; kết quả sensitivity analysis ở RP5 sẽ xác nhận hoặc điều chỉnh tỷ lệ này.")
    add_text_before(ref_523, "Tương tự, B ưu tiên Factual Accuracy hơn Grade-Level Appropriateness vì câu trả lời sai kiến thức gây rủi ro trực tiếp; C phân bổ cao hơn cho Robustness và Coherence vì hệ thống được đánh giá trong hội thoại nhiều lượt; D chia đều giữa Hallucination Control và Response Speed. Overall dùng 0.40A + 0.25B + 0.25C + 0.10D để đặt mục tiêu Socratic lên trước nhưng vẫn giữ Accuracy và Pedagogy là các điều kiện chất lượng chính. Các trọng số này không được chọn bằng locked test.")
    add_text_before(ref_523, "Căn cứ dữ liệu hiện có chỉ đủ để đặt giả thuyết, chưa đủ để khẳng định tối ưu. Pilot cũ có 9/10 lượt vi phạm A1, nên hỗ trợ việc coi A1 là hard constraint; pilot đó không được dùng để fit toàn bộ trọng số. Router test 60 case chỉ đo routing, không có raw response nên không được dùng để chọn điểm A/B/C/D.")

    # Explicit validation protocol, with status clearly separated from results.
    add_heading_before(ref_523, "5.2.2.2. Validation subset và độ đồng thuận người–máy")
    add_text_before(ref_523, "Validation subset dự kiến gồm N=30 conversation samples lấy từ development response set sau khi scope môn học được khóa, phân tầng theo môn và intent, mỗi conversation chọn một assistant response mục tiêu. Các mẫu này bị loại khỏi locked response test và không được dùng để báo cáo kết quả cuối.")
    add_text_before(ref_523, "Mỗi mẫu được chấm độc lập bởi một giáo viên/mentor có kinh nghiệm môn học, một researcher đã được hướng dẫn rubric và Gemini Judge với prompt version cố định. Human consensus là điểm trung vị hoặc nhãn đồng thuận của hai người chấm; Gemini không được xem điểm của người chấm trước khi nộp kết quả.")
    add_text_before(ref_523, "Độ đồng thuận được đo bằng weighted Cohen’s κ cho từng tiêu chí trên thang 0–5 và cho nhãn vi phạm A1 dạng nhị phân. Ngưỡng chấp nhận trước là κ≥0.60, tương ứng mức agreement khá trở lên. Nếu κ thấp hơn ngưỡng, nhóm phải rà soát rubric/prompt và báo cáo kết quả sensitivity thay vì coi Gemini Judge là ground truth. Hiện RP4 mới định nghĩa protocol; giá trị κ chưa được đo và sẽ báo cáo trong RP5.")

    # Add latency decomposition and the six planned scenarios before limitations.
    ref_lim = find_first(doc, "6. Giới hạn của phương pháp nghiên cứu")
    add_heading_before(ref_lim, "5.4.1. Latency decomposition và stability protocol")
    add_text_before(ref_lim, "Mỗi lượt End-to-End được ghi timestamp theo bốn mốc: request_start, router_decision, first_token và generation_end. Từ đó báo cáo router_latency_ms = router_decision − request_start; model_switch_latency_ms = thời gian chờ/khởi tạo model mới, bằng 0 nếu model không đổi; TTFT = first_token − request_start; generation_latency_ms = generation_end − first_token; end_to_end_latency_ms = generation_end − request_start.")
    add_text_before(ref_lim, "Router decision stability và response latency stability là hai khái niệm khác nhau. Decision stability là tỷ lệ lặp lại cùng signature (subject, intent, clarification, strategy). Latency stability dùng coefficient of variation: CV_latency = σ_latency / μ_latency. Với response text, có thể báo cáo thêm exact match hoặc mean pairwise Jaccard trên các lần replay; không dùng một chỉ số để thay thế chỉ số kia.")
    add_text_before(ref_lim, "Sáu kịch bản latency/switching được thiết kế cho RP5: (1) cold start khi chưa có model trong GPU; (2) lặp lại cùng subject/model; (3) chuyển General → specialist; (4) quay lại specialist/model cũ → General; (5) follow-up giữ nguyên subject; (6) câu mơ hồ dẫn đến clarification hoặc fallback General. Mỗi kịch bản chạy tối thiểu 3 repetitions trên subset cố định; RP4 chỉ mô tả protocol, không đưa số đo thực nghiệm.")

    # PhoBERT is not used by the current router implementation; remove orphan reference.
    remove_reference(doc, "[8] Nguyen, D. Q., & Nguyen, A. T.")

    # Repair cross-references after moving results to RP5.
    set_paragraph(
        find_first(doc, "Để đảm bảo tính so sánh công bằng với Version 2"),
        "Để đảm bảo tính so sánh công bằng với Version 2, nhóm tiến hành chuẩn hóa lại quy trình đánh giá baseline: giữ nguyên mô hình nền tảng của Version 1 (Qwen3-0.6B), trích xuất riêng phần phản hồi của AI trên tập pilot, sau đó chấm điểm bằng Gemini Judge và rubric được mô tả tại §5.2.2. Đây là protocol phương pháp; kết quả baseline được báo cáo ở RP5.",
    )
    set_paragraph(
        find_first(doc, "Chiến lược đánh giá tách hai câu hỏi"),
        "Chiến lược đánh giá tách hai câu hỏi: Router có chọn đúng model với chi phí hợp lý hay không (RP4), và model được chọn có tạo phản hồi đúng/Socratic hơn pooled model hay không (RP5). Cả ba Router mode dùng cùng locked routing set; RP4 chỉ định nghĩa protocol và metric, còn kết quả thực nghiệm được trình bày ở RP5.",
    )
    set_paragraph(
        find_first(doc, "Các công thức trên được áp dụng cho benchmark Router sơ bộ"),
        "Các công thức trong chương này định nghĩa cách tính metric cho Router và End-to-End. Report 4 không đưa kết quả benchmark; điểm chất lượng, Socratic, factuality và generation latency sẽ được tính sau khi có replay Pooled/Oracle/Hybrid và báo cáo trong RP5.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
