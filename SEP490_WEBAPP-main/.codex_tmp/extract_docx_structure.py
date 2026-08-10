import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    source = Path(sys.argv[1])
    doc = Document(source)
    paragraphs = []
    drawings = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": idx,
                    "style": paragraph.style.name if paragraph.style else "",
                    "text": text,
                }
            )
        blips = paragraph._p.xpath('.//a:blip')
        if blips:
            drawings.append(
                {
                    "paragraph_index": idx,
                    "text": paragraph.text.strip(),
                    "relationships": [
                        {
                            "id": blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'),
                            "target": str(doc.part.rels[blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')].target_ref),
                        }
                        for blip in blips
                    ],
                }
            )

    tables = []
    for table_idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_idx, "rows": rows})

    sections = []
    for idx, section in enumerate(doc.sections):
        sections.append(
            {
                "index": idx,
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
            }
        )

    print(
        json.dumps(
            {
                "source": str(source),
                "paragraph_count": len(doc.paragraphs),
                "nonempty_paragraphs": paragraphs,
                "table_count": len(doc.tables),
                "tables": tables,
                "inline_shape_count": len(doc.inline_shapes),
                "drawings": drawings,
                "sections": sections,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
