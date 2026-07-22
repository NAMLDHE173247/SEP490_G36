from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(r'D:\Giai_thich_V1_V2_va_cau_hoi_nghien_cuu.docx')
BLUE='1F4E79'; NAVY='17365D'; LIGHT='EAF2F8'; GREEN='E2F0D9'; ORANGE='FCE4D6'; GRAY='F2F2F2'

def cell(c,txt,fill=LIGHT,color='000000',bold=False,size=10):
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); c._tc.get_or_add_tcPr().append(shd)
    r=p.add_run(txt); r.bold=bold; r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)

def h(d,text,level=1):
    p=d.add_paragraph(style=f'Heading {level}'); r=p.add_run(text); r.font.name='Arial'; r.font.color.rgb=RGBColor.from_string(BLUE); return p
def p(d,text):
    q=d.add_paragraph(); q.paragraph_format.space_after=Pt(5); q.paragraph_format.line_spacing=1.1; r=q.add_run(text); r.font.name='Arial'; r.font.size=Pt(11); return q
def bullets(d,items):
    for x in items:
        q=d.add_paragraph(style='List Bullet'); q.paragraph_format.space_after=Pt(2); r=q.add_run(x); r.font.name='Arial'; r.font.size=Pt(10.5)
def flow(d,nodes,fills):
    t=d.add_table(rows=1,cols=len(nodes)*2-1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,n in enumerate(nodes):
        cell(t.cell(0,i*2),n,fills[i],NAVY,True,9)
        if i<len(nodes)-1: cell(t.cell(0,i*2+1),'→','FFFFFF',BLUE,True,16)

def main():
    d=Document(); s=d.sections[0]; s.top_margin=Inches(.65); s.bottom_margin=Inches(.65); s.left_margin=Inches(.7); s.right_margin=Inches(.7)
    d.styles['Normal'].font.name='Arial'; d.styles['Normal'].font.size=Pt(11)
    for nm,sz in [('Heading 1',18),('Heading 2',14)]: d.styles[nm].font.name='Arial'; d.styles[nm].font.size=Pt(sz)
    title=d.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=title.add_run('GIẢI THÍCH V1 – V2 VÀ CÂU HỎI NGHIÊN CỨU'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(22); r.font.color.rgb=RGBColor.from_string(NAVY)
    sub=d.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=sub.add_run('Bản nói ngắn cho buổi trao đổi với mentor'); r.font.name='Arial'; r.font.size=Pt(12)
    h(d,'1. Nói một câu để mentor hiểu ngay',2)
    p(d,'“V1 dùng một model chung để trả lời mọi câu hỏi. V2 vẫn giữ mục tiêu gia sư Socratic, nhưng tách model theo môn và dùng Router để quyết định câu nào nên vào model nào. Nghiên cứu của nhóm là kiểm tra việc tách này có thực sự đáng làm hay không.”')
    h(d,'2. V1 thực sự là gì?',2)
    flow(d,['Dữ liệu\nhội thoại chung','Một shared\nQwen tutor','Trả lời mọi\ncâu hỏi'],[LIGHT,ORANGE,LIGHT])
    bullets(d,['V1 có một model chung: không có Math model, Physics model hay Chemistry model riêng.','Ưu điểm: đơn giản, ít thành phần, dễ chạy.','Vấn đề: model chung phải học tất cả môn; khi trả lời sai rất khó biết lỗi do dữ liệu, do model hay do không có chuyên môn theo môn.','V1 là “vấn đề/điểm xuất phát” của nghiên cứu, không phải POOLED GENERAL adapter trong thí nghiệm V2.'])
    h(d,'3. V2 khác V1 ở đúng chỗ nào?',2)
    flow(d,['Câu hỏi\n+ history','Router\nquyết định','Model phù hợp\nđược gọi','Phản hồi\nSocratic'],[LIGHT,GREEN,ORANGE,LIGHT])
    p(d,'V2 có hai loại model khác nhau. Đây là điểm quan trọng nhất:')
    t=d.add_table(rows=1,cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(['Thành phần','Nó làm gì?','Nó không làm gì?']): cell(t.cell(0,i),x,BLUE,'FFFFFF',True,9)
    for row in [('Router','Đọc câu hỏi + history, chọn route','Không giải bài cho học sinh'),('GENERAL / MATH / PHYSICS / CHEMISTRY','Sinh phản hồi theo phong cách Socratic','Không tự chọn chính nó'),('Gemini Judge','Chấm replay sau thí nghiệm','Không tham gia trả lời học sinh')]:
        cs=t.add_row().cells
        for i,x in enumerate(row): cell(cs[i],x,'FFFFFF','000000',False,9)
    d.add_page_break()
    h(d,'4. Luồng V2 từng bước',1)
    for n,txt in [
      ('Bước 1 – Nhận input','Học sinh hỏi một câu; hệ thống kèm history gần nhất nếu đây là follow-up.'),
      ('Bước 2 – Rule Router','Rule kiểm tra keyword/regex và previous subject. Case rất rõ có thể được chọn ngay, nên nhanh và không mất LLM call.'),
      ('Bước 3 – Confidence gate','Nếu rule không đủ chắc, câu mơ hồ, viết sai, liên môn hoặc follow-up thì chuyển sang LLM-as-Router.'),
      ('Bước 4 – LLM-as-Router','LLM chỉ trả decision có cấu trúc: subject, intent, need_clarification và target model. Nó không giải bài.'),
      ('Bước 5 – Gọi response model','Hệ thống gọi GENERAL V2 hoặc specialist đúng môn. Đây là các LoRA adapter đã train.'),
      ('Bước 6 – Log và chấm','Lưu route, latency, LLM calls và response. Chỉ sau đó Model Eval/Gemini Judge mới chấm chất lượng phản hồi.')]:
        p(d,n+' : '+txt)
    h(d,'Ví dụ để nói',2)
    p(d,'Học sinh: “Sao chỗ này lại nhân 1/2?”')
    bullets(d,['Không có history: Router không biết “chỗ này” là gì → ask clarification.','Có history “em đang học rơi tự do, s = 1/2gt²”: Router nhận ra PHYSICS → gọi Physics specialist.','Như vậy V2 không mạnh vì “nhiều model” một cách tự động; nó chỉ có ích khi Router dùng đúng history để gọi đúng model.'])
    h(d,'5. Câu hỏi nghiên cứu thật sự là gì?',1)
    p(d,'Câu hỏi chính: “Hybrid Router có đưa câu hỏi STEM tiếng Việt đến đúng specialist và tạo phản hồi tốt hơn một model chung, với độ trễ/chi phí chấp nhận được hay không?”')
    bullets(d,['RQ1: Hybrid có tăng Subject Accuracy và Macro-F1 so với Rule hoặc LLM Router đơn lẻ không?','RQ2: History/previous subject có làm follow-up route tốt hơn không?','RQ3: Khi Router chọn specialist, phản hồi HYBRID có gần Oracle và tốt hơn GENERAL V2 về Socratic/factuality không?','RQ4: Lợi ích chất lượng có đổi lại bằng latency và LLM-call rate bao nhiêu?'])
    d.add_page_break()
    h(d,'6. Làm sao kiểm chứng mà không nói quá?',1)
    h(d,'RP4: kiểm tra Router',2)
    p(d,'60 câu hold-out (20 Math, 20 Physics, 20 Chemistry) → RULE / LLM / HYBRID. Đo Subject Accuracy, Macro-F1, Exact Match, Wrong-route, latency, LLM-call rate và ablation history/threshold.')
    h(d,'RP5: kiểm tra phản hồi',2)
    flow(d,['Cùng 1\nprompt test','B: GENERAL\nV2','C: ORACLE\nspecialist','D: HYBRID\nrouter'],[LIGHT,ORANGE,GREEN,LIGHT])
    p(d,'B là baseline V2 (GENERAL adapter trả lời mọi câu), C là upper bound vì biết sẵn gold subject, D là hệ thống thật. Gemini Judge chấm Socratic, factuality, overall; đây là proxy, không phải chứng minh học sinh tiến bộ.')
    h(d,'Công thức nói ngắn gọn',2)
    bullets(d,['Accuracy = số route đúng / tổng số case.','Macro-F1 = trung bình F1 của từng môn, tránh kết quả đẹp giả do lệch môn.','LLM-call rate = số case gọi LLM Router / tổng case.','Hybrid chỉ hợp lý khi chất lượng D gần C nhưng latency/LLM calls thấp hơn việc gọi LLM Router cho mọi case.'])
    h(d,'7. 4 câu hỏi để hỏi mentor',1)
    bullets(d,['Scope Hybrid routing + LoRA specialists + offline evaluation có đủ cho capstone không?','Có cần human review nhỏ để hỗ trợ Gemini Judge không?','Nên lấy Router accuracy hay B/C/D quality trade-off làm claim chính?','Tên đề tài có nên thu hẹp theo routing để khớp với thực nghiệm không?'])
    h(d,'8. Câu chốt buổi họp',2)
    p(d,'“Nhóm không khẳng định V2 tốt hơn trước. Nhóm đã có V2 để kiểm chứng; sau khi mentor xác nhận scope, nhóm sẽ chạy RP4 và RP5 trên cùng test set, báo cáo cả kết quả tốt lẫn lỗi và giới hạn.”')
    d.save(OUT); print(OUT)

if __name__=='__main__': main()
