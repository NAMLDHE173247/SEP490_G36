from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"D:\Brief_Nghien_cuu_Gui_Mentor_30_phut.docx")
BLUE = "1F4E79"; LIGHT = "EAF2F8"; NAVY = "17365D"; GREEN = "E2F0D9"; ORANGE = "FCE4D6"; GRAY = "F2F2F2"

def set_cell(cell, text, fill=LIGHT, color="000000", bold=False, size=10):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = 'Arial'; r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')

def borderless(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'nil'); borders.append(el)
    tblPr.append(borders)

def heading(doc, text, level=1):
    p=doc.add_paragraph(); p.style=f'Heading {level}'
    r=p.add_run(text); r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); r.font.color.rgb=RGBColor.from_string(BLUE)
    return p

def body(doc, text, bold_prefix=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.1
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); r.bold=True; r.font.name='Arial'; r.font.size=Pt(10.5)
        r=p.add_run(text[len(bold_prefix):]); r.font.name='Arial'; r.font.size=Pt(10.5)
    else:
        r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(10.5)
    return p

def bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2)
        r=p.add_run(item); r.font.name='Arial'; r.font.size=Pt(10.5)

def pipeline(doc, nodes, fills):
    t=doc.add_table(rows=1, cols=len(nodes)*2-1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borderless(t)
    for i,node in enumerate(nodes):
        cell=t.cell(0,i*2); cell.width=Inches(1.02); set_cell(cell,node,fills[i],NAVY,True,8.5)
        if i<len(nodes)-1:
            a=t.cell(0,i*2+1); a.width=Inches(.22); set_cell(a,'→','FFFFFF',BLUE,True,15)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def comparison(doc):
    t=doc.add_table(rows=1, cols=4); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    headers=['Điều kiện','Model được gọi','Vai trò nghiên cứu','Kết quả cần so sánh']
    for i,h in enumerate(headers): set_cell(t.cell(0,i),h,BLUE,'FFFFFF',True,9)
    rows=[
      ('B. POOLED','GENERAL cho mọi câu','Baseline hệ thống đơn','Quality / latency'),
      ('C. ORACLE','Specialist theo gold subject','Upper bound model selection','Quality tối đa khả dĩ'),
      ('D. HYBRID','Router tự chọn GENERAL/specialist','Phương pháp đề xuất','Quality–latency trade-off'),
    ]
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell(cells[i],val,'FFFFFF','000000',False,9)

def main():
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    styles=d.styles; styles['Normal'].font.name='Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); styles['Normal'].font.size=Pt(10.5)
    for s,size in [('Heading 1',17),('Heading 2',13)]: styles[s].font.name='Arial'; styles[s]._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); styles[s].font.size=Pt(size)

    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    r=p.add_run('BRIEF NGHIÊN CỨU GỬI MENTOR'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
    r=p.add_run('Hybrid Multi-Model Routing cho AI Tutor STEM tiếng Việt'); r.font.name='Arial'; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(BLUE)
    t=d.add_table(rows=1,cols=1); borderless(t); set_cell(t.cell(0,0),'Mục tiêu buổi họp: xác nhận hướng nghiên cứu, phạm vi và thiết kế kiểm chứng trước khi chốt RP4–RP5.',GREEN,NAVY,True,11)
    heading(d,'1. Điều cần mentor nắm trong 2 phút',2)
    bullets(d,[
      'Nhóm không nghiên cứu UI; câu hỏi chính là router nào chọn model phù hợp nhất cho tutor STEM tiếng Việt.',
      'V1 là baseline đơn/định tuyến đơn giản. V2 là Hybrid: rule xử lý case rõ, LLM Router hỗ trợ case khó, sau đó gọi GENERAL hoặc specialist.',
      'Đã train 4 LoRA adapters trên cùng base model: GENERAL 600 mẫu; MATH, PHYSICS, CHEMISTRY 200 mẫu/môn. Kết quả chưa được kết luận trước benchmark.',
      'Cần mentor chốt: scope offline evaluation có đủ cho capstone không; và cách diễn giải kết quả synthetic dataset/Judge.'
    ])
    heading(d,'2. RP1 → RP4: một mạch nghiên cứu',2)
    comparison_t=d.add_table(rows=1,cols=3); comparison_t.style='Table Grid'
    for i,h in enumerate(['Report','Đã xác định','Đưa vào brief này']): set_cell(comparison_t.cell(0,i),h,BLUE,'FFFFFF',True,9)
    for row in [
      ('RP1','Bối cảnh Flipped Classroom, tutor Socratic, vấn đề đa môn','Problem + significance'),
      ('RP2','Research questions, feasibility, roadmap','Scope + hypotheses'),
      ('RP3','SLM/LoRA, routing, Socratic literature','Gap + baseline rationale'),
      ('RP4','Thuật toán, data split, metrics, experiment','Pipeline + formulas + evaluation'),
    ]:
      cells=comparison_t.add_row().cells
      for i,v in enumerate(row): set_cell(cells[i],v,'FFFFFF','000000',False,9)
    d.add_page_break()

    heading(d,'3. Pipeline V1 – baseline cần vượt qua',1)
    pipeline(d,['Câu hỏi\nhọc sinh','General tutor\n/ rule đơn','Phản hồi'],[LIGHT,ORANGE,LIGHT])
    body(d,'Ý nghĩa: V1 đơn giản, dễ vận hành và là baseline công bằng. Điểm cần kiểm chứng là liệu V1 có bị giảm chất lượng ở câu phụ thuộc history, mơ hồ hoặc cần chuyên môn theo môn hay không.')
    heading(d,'4. Pipeline V2 – Hybrid Multi-Model Routing',1)
    pipeline(d,['Question\n+ history','Rule\nRouter','Confidence\ngate','LLM-as-\nRouter','Model\nselection','Socratic\nresponse'],[LIGHT,GREEN,ORANGE,LIGHT,GREEN,LIGHT])
    body(d,'Luồng quyết định: nếu rule đủ tự tin thì đi thẳng sang model; nếu low-confidence, mơ hồ hoặc follow-up thì gọi LLM-as-Router. Model selection có thể chọn GENERAL, MATH, PHYSICS, CHEMISTRY hoặc ask clarification.')
    t=d.add_table(rows=1,cols=2); t.style='Table Grid'
    set_cell(t.cell(0,0),'Case',[BLUE][0],'FFFFFF',True,9); set_cell(t.cell(0,1),'Quy tắc V2',[BLUE][0],'FFFFFF',True,9)
    for row in [('Câu rõ','Rule → specialist trực tiếp; giảm LLM calls/latency'),('Mơ hồ / follow-up','History + LLM Router → quyết định hoặc hỏi lại'),('Router/LLM lỗi','Fallback có log; không giấu lỗi khi phân tích')]:
      cells=t.add_row().cells; set_cell(cells[0],row[0],GRAY,'000000',True,9); set_cell(cells[1],row[1],'FFFFFF','000000',False,9)
    heading(d,'5. Research gap và đóng góp thực tế',2)
    body(d,'Khoảng trống nhóm không tuyên bố là “chưa ai làm router”. Khoảng trống cụ thể là thiếu một đánh giá thực nghiệm có thể tái lập cho AI Tutor STEM tiếng Việt, nơi routing phải xét subject, intent, context, clarification và chất lượng phản hồi Socratic.')
    bullets(d,['Đóng góp 1: Hybrid Router + logging/ablation để biết context và threshold có đóng góp gì.','Đóng góp 2: so sánh end-to-end GENERAL vs Oracle specialist vs Hybrid trên cùng prompt.','Đóng góp 3: công bố rõ giới hạn của dữ liệu synthetic và judge-based evaluation.'])
    d.add_page_break()

    heading(d,'6. Thiết kế kiểm chứng và công thức',1)
    heading(d,'RP4 – Router Benchmark',2)
    body(d,'Tập hold-out: 60 cases, cân bằng 20 MATH / 20 PHYSICS / 20 CHEMISTRY; có viết tắt, lỗi chính tả, câu mơ hồ và follow-up. So sánh RULE, LLM, HYBRID; 3 repeats để đo Stability; ablation: no-history, no-previous-subject, strict, lenient.')
    body(d,'Subject Accuracy = (1/N) Σ I(ŷ_subject = y_subject). Macro-F1 = (1/K) Σ F1_k. Exact Match = (1/N) Σ I(ŷ_subject, ŷ_intent, ŷ_clarification đều đúng).')
    body(d,'Latency trung bình: L̄ = (1/N) Σ L_i. LLM-call rate = (1/N) Σ I(LLM Router được gọi). Hai metric này phù hợp với Hybrid vì nó phải gọi LLM có chọn lọc, không chỉ tối đa accuracy.')
    heading(d,'RP5 – End-to-End Response Evaluation',2)
    comparison(d)
    body(d,'Quality_j = (1/N) Σ score_j, với j ∈ {Socratic, factuality, overall}. Cùng một prompt được dùng cho B/C/D để tách: năng lực GENERAL, upper bound specialist và lỗi truyền từ Router.')
    heading(d,'Vì sao chọn Hybrid thay vì RL?',2)
    body(d,'Tiêu chí lý thuyết: a* = argmax_a [w_Q·Q(a) − w_L·L(a)/L_max − w_C·C(a)/C_max]. V2 chỉ hiện thực hóa trực giác này qua confidence threshold, không tuyên bố đã train chính sách RL hay tối ưu trọng số. RL/RCT là future work vì cần reward ổn định, dữ liệu thật và quy trình đạo đức.')
    d.add_page_break()

    heading(d,'7. Trạng thái, rủi ro và quyết định cần mentor',1)
    t=d.add_table(rows=1,cols=3); t.style='Table Grid'
    for i,h in enumerate(['Hạng mục','Trạng thái','Việc cần chốt']): set_cell(t.cell(0,i),h,BLUE,'FFFFFF',True,9)
    for row in [
      ('Router Hybrid / RP4','Đã có code + test set','Chạy benchmark, không tune trên hold-out'),
      ('4 LoRA adapters','Đã train','Export Hub + Registry Active hfRepoId'),
      ('RP5 / Model Eval','Đã có luồng replay','Chạy B/C/D + Judge rubric'),
      ('Độ tin cậy','Synthetic 60 cases, judge proxy','Nêu limitation; cân nhắc human review mẫu nhỏ'),
    ]:
      cells=t.add_row().cells
      for i,v in enumerate(row): set_cell(cells[i],v,'FFFFFF','000000',False,9)
    heading(d,'Câu hỏi đề nghị mentor phản hồi',2)
    bullets(d,[
      'Có chấp nhận scope: Hybrid routing + LoRA specialists + offline evaluation cho capstone, không làm RL/RCT?',
      'Có cần bổ sung human review/giáo viên trên một mẫu nhỏ để hỗ trợ Gemini Judge không?',
      'Ưu tiên kết luận nào: Router Macro-F1, chênh lệch quality B/C/D, hay trade-off latency/LLM-call rate?',
      'Tên đề tài nên giữ “Fine-tuning AI for Flipped Classroom” hay thu hẹp để khớp thực nghiệm routing?' 
    ])
    heading(d,'Lộ trình nói 30 phút',2)
    bullets(d,['0–4 phút: vấn đề và RP1–RP3.','4–12 phút: Pipeline V1 và V2.','12–20 phút: thiết kế RP4/RP5 và công thức.','20–25 phút: trạng thái train/Registry, rủi ro.','25–30 phút: 4 câu hỏi cần mentor quyết định.'])
    d.save(OUT); print(OUT)

if __name__=='__main__': main()
