from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


ROOT = Path(r"D:\Sep_G36\SEP490_G36\SEP490_WEBAPP-main")
OUT = ROOT / "docs" / "Bao_cao_co_so_lua_chon_pretrained_model_theo_mon_hoc.docx"
ASSET_DIR = ROOT / ".codex_tmp" / "mentor_model_report_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)
PIPELINE_IMG = ASSET_DIR / "pipeline_lua_chon_model.png"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MID_BLUE = "4C78A8"
PALE_BLUE = "E8EEF5"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "9A6A00"
RED = "9B1C1C"
GREEN = "2F6B4F"
GRAY = "5A6673"
WHITE = "FFFFFF"
BLACK = "111111"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color=BLUE, size=10, space=7):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    rfonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Calibri")
    rpr.append(rfonts)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc, text="", bold_lead=None, italic=False, align=None, after=6):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)
    return p


def create_numbering_id(doc):
    numbering = doc.part.numbering_part.element
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = (max(num_ids) + 1) if num_ids else 1
    # Reuse Word's built-in List Number abstract definition. Creating new
    # abstract definitions can make some Word/LibreOffice versions pause on
    # repair; a new concrete num with a start override is schema-safe.
    abstract_id = 7
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def add_number(doc, text, num_id):
    p = doc.add_paragraph(style="Normal")
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(nid)
    ppr.insert(0, num_pr)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc, label, text, fill=CALLOUT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], 120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), accent)
    tc_borders.append(left)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def style_table(table, header=True, font_size=9.2, header_fill=LIGHT, center_cols=None):
    center_cols = set(center_cols or [])
    table.style = "Table Grid"
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if header and r_idx == 0:
                set_cell_shading(cell, header_fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                if c_idx in center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(header and r_idx == 0), color=NAVY if header and r_idx == 0 else BLACK)
    if header:
        set_repeat_table_header(table.rows[0])


def add_table_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    keep_with_next(p)
    r = p.add_run(text)
    set_run_font(r, size=9.5, bold=True, color=DARK_BLUE)
    return p


def draw_pipeline(path):
    width, height = 1800, 470
    img = Image.new("RGB", (width, height), "#" + WHITE)
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    f_title = ImageFont.truetype(bold_path, 36)
    f_box = ImageFont.truetype(bold_path, 26)
    f_sub = ImageFont.truetype(font_path, 21)
    draw.text((60, 28), "PIPELINE LỰA CHỌN MODEL THEO MÔN HỌC", font=f_title, fill="#0B2545")
    labels = [
        ("1. Ứng viên", "0.5B-1.7B\nopen-weight"),
        ("2. Base test", "Chưa fine-tune\ncùng prompt"),
        ("3. LoRA/SFT", "Cùng data, seed,\nhyperparameters"),
        ("4. Locked test", "Toán - Anh - Sử\nkhông rò rỉ"),
        ("5. Kiểm định", "CI 95% + paired\nsignificance"),
        ("6. Triển khai", "Registry → Router\n→ E2E RP5"),
    ]
    x0, y, box_w, box_h, gap = 55, 135, 245, 215, 45
    for i, (title, sub) in enumerate(labels):
        x = x0 + i * (box_w + gap)
        fill = "#E8EEF5" if i < 5 else "#DDEFE5"
        outline = "#2E74B5" if i < 5 else "#2F6B4F"
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=22, fill=fill, outline=outline, width=4)
        bbox = draw.multiline_textbbox((0, 0), title, font=f_box, align="center")
        tw = bbox[2] - bbox[0]
        draw.text((x + (box_w - tw) / 2, y + 34), title, font=f_box, fill="#0B2545")
        bbox2 = draw.multiline_textbbox((0, 0), sub, font=f_sub, align="center", spacing=7)
        sw = bbox2[2] - bbox2[0]
        draw.multiline_text((x + (box_w - sw) / 2, y + 104), sub, font=f_sub, fill="#303B46", align="center", spacing=7)
        if i < len(labels) - 1:
            ax1, ax2, ay = x + box_w + 7, x + box_w + gap - 7, y + box_h / 2
            draw.line((ax1, ay, ax2, ay), fill="#4C78A8", width=7)
            draw.polygon([(ax2, ay), (ax2 - 16, ay - 12), (ax2 - 16, ay + 12)], fill="#4C78A8")
    draw.text((60, 405), "Nguyên tắc: không dùng kết quả benchmark của nhà sản xuất để thay thế thử nghiệm trên dữ liệu dự án.", font=f_sub, fill="#5A6673")
    img.save(path, quality=95)


def setup_document():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        st = styles[list_name]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167

    if "Reference" not in [s.name for s in styles]:
        ref = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        ref = styles["Reference"]
    ref.font.name = "Calibri"
    ref.font.size = Pt(9.5)
    ref._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    ref._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    ref._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    ref.paragraph_format.left_indent = Inches(0.28)
    ref.paragraph_format.first_line_indent = Inches(-0.28)
    ref.paragraph_format.space_after = Pt(5)
    ref.paragraph_format.line_spacing = 1.05

    header = sec.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("SEP490 | RESEARCH BRIEF")
    set_run_font(r, size=8.5, bold=True, color=GRAY)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Hybrid Multi-SLM Router  |  Trang ")
    set_run_font(r, size=8.5, color=GRAY)
    add_page_field(fp)
    return doc


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("BÁO CÁO CƠ SỞ KHOA HỌC")
    set_run_font(r, size=9.5, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("LỰA CHỌN PRETRAINED MODEL\nTHEO MÔN HỌC")
    set_run_font(r, size=24, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Hybrid Multi-SLM Router cho Toán, Tiếng Anh và Lịch sử")
    set_run_font(r, size=14, color=DARK_BLUE)

    rows = [
        ("Mục đích:", "Báo cáo mentor - rà soát cơ sở chọn model và thiết kế thí nghiệm"),
        ("Phạm vi:", "Pretrained model, hallucination, chất lượng sư phạm và chi phí triển khai"),
        ("Ngày cập nhật:", "20/07/2026"),
        ("Trạng thái:", "Đề xuất phương pháp; chưa thay thế kết quả thực nghiệm RP5"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(label + " ")
        set_run_font(r1, bold=True, color=NAVY)
        r2 = p.add_run(value)
        set_run_font(r2)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    paragraph_border_bottom(rule, color=BLUE, size=12, space=5)


def build_document():
    draw_pipeline(PIPELINE_IMG)
    doc = setup_document()
    add_title_block(doc)

    add_callout(
        doc,
        "KẾT LUẬN NGẮN",
        "Lựa chọn Qwen2.5-0.5B hiện có cơ sở kỹ thuật (nhỏ, open-weight, chạy được LoRA trên phần cứng nhóm) nhưng chưa có đủ cơ sở khoa học để gọi là model tốt nhất cho từng môn. Cách bảo vệ đúng là coi đây là baseline bị ràng buộc tài nguyên, rồi lựa chọn model bằng thí nghiệm đối chứng trên cùng dữ liệu khóa.",
        fill=PALE_BLUE,
    )

    doc.add_heading("1. Câu trả lời trực tiếp cho yêu cầu của mentor", level=1)
    add_body(doc, "Câu hỏi của mentor gồm hai phần: (i) có công trình uy tín nào đánh giá hoặc lựa chọn model theo môn học hay không; và (ii) có cơ sở nào để đánh giá hallucination theo từng model hay không.")
    num_id = create_numbering_id(doc)
    add_number(doc, "Có nhiều benchmark uy tín đánh giá theo môn, ngôn ngữ và bối cảnh giáo dục; tuy nhiên không có một bài báo duy nhất cho phép kết luận sẵn rằng một model cụ thể luôn tối ưu cho Toán, Tiếng Anh hoặc Lịch sử.", num_id)
    add_number(doc, "Bằng chứng mạnh nhất cho thấy năng lực phụ thuộc vào môn học, loại nhiệm vụ, ngôn ngữ, prompt và mục tiêu sư phạm. Vì vậy, việc chọn model phải là một thí nghiệm độc lập trong dự án, không phải quyết định dựa trên bảng quảng bá của nhà phát triển.", num_id)
    add_number(doc, "Hallucination phải được đo bằng dữ liệu riêng của từng môn. Lịch sử ưu tiên factual precision và unsupported-claim rate; Toán ưu tiên độ đúng đáp án và độ đúng từng bước; Tiếng Anh ưu tiên độ đúng ngôn ngữ, tính nhất quán và chất lượng phản hồi sư phạm.", num_id)
    add_number(doc, "Nguồn chính của báo cáo là bài đã phản biện tại ACL, EMNLP, NeurIPS, ICML, AAAI và tạp chí ZDM. Technical report và arXiv chỉ dùng để lập danh sách model ứng viên hoặc future work, không dùng làm bằng chứng độc lập để tuyên bố model tốt nhất.", num_id)

    doc.add_heading("2. Phân tầng chất lượng nguồn", level=1)
    add_body(doc, "Không nên trộn tất cả nguồn vào cùng một mức. SCImago phù hợp chủ yếu với tạp chí; với hội nghị khoa học máy tính, cần ghi rõ venue và trạng thái peer review thay vì gán SJR của một ấn phẩm khác cho bài hội nghị.")
    add_table_caption(doc, "Bảng 1. Quy tắc ưu tiên nguồn dùng trong báo cáo")
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    for i, t in enumerate(["Mức", "Loại nguồn", "Cách sử dụng", "Nguồn tiêu biểu"]):
        hdr[i].text = t
    rows = [
        ("A", "Peer-reviewed", "Làm cơ sở cho lập luận, metric và thiết kế thí nghiệm.", "ACL, EMNLP, NeurIPS, ICML, AAAI, ZDM"),
        ("B", "Technical report", "Dùng mô tả kiến trúc, kích thước và số liệu do tác giả model công bố; phải ghi rõ xung đột lợi ích.", "Qwen2.5, Qwen2.5-Math"),
        ("C", "Preprint/arXiv", "Chỉ dùng gợi ý future work hoặc phương pháp mới cần kiểm chứng.", "OpenLearnLM, K12-KGraph, Theorem-SFT"),
        ("Loại", "Blog/leaderboard không phương pháp", "Không dùng làm bằng chứng chính; chỉ dùng tìm manh mối rồi truy về nguồn gốc.", "Medium, trang tổng hợp benchmark, bài quảng bá"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [900, 1700, 4100, 2660])
    style_table(table, font_size=9.0, center_cols=[0])

    doc.add_heading("3. Tổng hợp các công trình có giá trị trực tiếp", level=1)
    add_table_caption(doc, "Bảng 2. Nguồn cốt lõi, xếp theo mức độ liên quan và độ tin cậy")
    table = doc.add_table(rows=1, cols=4)
    for i, t in enumerate(["Công trình", "Venue / năm", "Bằng chứng chính", "Ý nghĩa cho dự án"]):
        table.rows[0].cells[i].text = t
    core = [
        ("VMLU Benchmarks", "ACL 2025", "Bốn tập đánh giá tiếng Việt: kiến thức chung, đọc hiểu, suy luận và hội thoại; so sánh Llama-3, Qwen2.5 và GPT-4.", "Bắt buộc kiểm tra tiếng Việt; điểm tiếng Anh hoặc tiếng Trung không thể thay thế."),
        ("K-12EduBench", "AAAI 2026", "2.640 câu khách quan + 619 câu tự luận, 9 môn; đo tri thức, giải quyết vấn đề và nhận thức mục tiêu giáo dục.", "Chọn model theo nhiều năng lực, không chỉ answer accuracy."),
        ("EduBench", "ACL 2026", ">4.000 bối cảnh, 9 kịch bản và 12 chiều đánh giá; có human annotation.", "Củng cố đánh giá đa chiều và kiểm chứng LLM judge bằng người."),
        ("MMLU-Pro", "NeurIPS 2024", "Câu hỏi khó hơn, 10 lựa chọn; điểm giảm 16-33% so với MMLU và ổn định hơn trước biến thể prompt.", "Dùng nhiệm vụ đủ khó; tránh benchmark bão hòa hoặc quá dễ."),
        ("E-EVAL", "Findings ACL 2024", "4.351 câu hỏi K-12 tiếng Trung, 23 môn; cho thấy ảnh hưởng rõ của ngôn ngữ và văn hóa.", "Là bằng chứng gián tiếp rằng Lịch sử phải đánh giá trong ngữ cảnh bản địa."),
        ("MathTutorBench", "EMNLP 2025", "Khả năng giải bài tốt không tự động đồng nghĩa với khả năng dạy tốt; hội thoại dài khó hơn.", "Tách điểm correctness khỏi điểm Socratic/pedagogy."),
        ("No one-size-fits-all", "ZDM 2026", "2.880 lời giải; model chi phối chất lượng nội dung, prompt chi phối quá trình và sư phạm; không có cấu hình thắng mọi chiều.", "Phải cố định prompt khi so model và báo cáo trade-off."),
        ("FActScore", "EMNLP 2023", "Tách câu trả lời thành atomic facts và tính tỷ lệ fact được nguồn đáng tin hỗ trợ.", "Metric chính cho Lịch sử và câu trả lời dài."),
        ("HaluEval / TruthfulQA", "EMNLP 2023 / ACL 2022", "Đo khả năng phát hiện hoặc tránh nội dung sai, cùng xu hướng lặp lại quan niệm sai phổ biến.", "Bổ sung false-premise và abstention cases; không suy hallucination từ fluency."),
        ("SciBench", "ICML 2024", "Toán, Hóa, Lý cấp đại học; model tốt nhất chỉ 43,22%; không prompt nào thắng mọi kỹ năng.", "Cần error taxonomy cho Toán, không chỉ exact answer."),
    ]
    for row in core:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [1700, 1450, 3150, 3060])
    style_table(table, font_size=8.4)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Nhận định: ")
    set_run_font(r, size=9.3, bold=True, color=DARK_BLUE)
    r = p.add_run("Các công trình trên hỗ trợ phương pháp lựa chọn theo môn, nhưng không trực tiếp chứng minh Qwen2.5-0.5B là lựa chọn tối ưu cho dữ liệu của nhóm.")
    set_run_font(r, size=9.3, italic=True, color=GRAY)

    doc.add_heading("4. Điều có thể và không thể kết luận từ literature", level=1)
    doc.add_heading("4.1. Có thể kết luận", level=2)
    for item in [
        "Năng lực model thay đổi đáng kể theo môn học, ngôn ngữ, loại câu hỏi và prompt.",
        "Năng lực giải đúng và năng lực dạy theo kiểu Socratic là hai mục tiêu khác nhau; cần đo riêng.",
        "Với tiếng Việt, phải có test set tiếng Việt hoặc song ngữ phù hợp; không thể dùng điểm benchmark tiếng Anh để đại diện.",
        "Hallucination là hiện tượng phụ thuộc nhiệm vụ. Cần định nghĩa claim có thể kiểm chứng, nguồn chuẩn và loại lỗi theo từng môn.",
        "Model nhỏ có thể phù hợp triển khai nhưng cần chứng minh non-inferiority về chất lượng trước khi chọn vì latency/VRAM.",
    ]:
        add_bullet(doc, item)
    doc.add_heading("4.2. Không được kết luận", level=2)
    for item in [
        "Không được đổi accuracy của một bộ phát hiện hallucination thành 'tỷ lệ model không hallucinate'. Đây là hai đại lượng khác nhau.",
        "Không được lấy bảng do nhà phát triển model công bố làm bằng chứng duy nhất để chọn model cho học sinh Việt Nam.",
        "Không được tuyên bố RAG luôn giảm hallucination theo một tỷ lệ cố định; hiệu quả phụ thuộc retrieval, corpus và nhiệm vụ.",
        "Không được gộp correctness, Socratic, latency và chi phí bằng trọng số tự đặt rồi gọi model có tổng điểm cao nhất là tối ưu.",
        "Không được dùng 60 cases Router Benchmark để lựa chọn answer model; bộ đó chỉ kiểm tra quyết định định tuyến.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("5. Đánh giá lựa chọn Qwen2.5 hiện tại", level=1)
    add_body(doc, "Qwen2.5-0.5B-Instruct là một lựa chọn khởi đầu hợp lý về engineering: open-weight, kích thước nhỏ, có thể LoRA và phù hợp tài nguyên GPU hạn chế. Tuy nhiên, các phản hồi smoke test của nhóm từng có dấu hiệu lẫn ngôn ngữ, câu gợi mở vô nghĩa và thiếu công thức. Đây là quan sát định tính để tạo hypothesis, chưa phải kết quả RP5 cho đến khi tái lập bằng batch test khóa.")
    add_table_caption(doc, "Bảng 3. Cách diễn đạt lựa chọn model trong báo cáo")
    table = doc.add_table(rows=1, cols=3)
    for i, t in enumerate(["Nội dung", "Cách viết chưa đủ", "Cách viết nên dùng"]):
        table.rows[0].cells[i].text = t
    rows = [
        ("Lý do chọn", "Qwen tốt nhất cho ba môn.", "Qwen2.5-0.5B là baseline bị ràng buộc tài nguyên và sẽ được so sánh có kiểm soát."),
        ("Benchmark hãng", "Điểm MATH cao nên chắc chắn dạy Toán tốt.", "Technical report cho thấy tiềm năng giải toán; MathTutorBench yêu cầu kiểm tra sư phạm riêng."),
        ("Kết luận", "Model đã train nên giữ nguyên.", "Registry chỉ xác nhận khả năng triển khai; không xác nhận chất lượng học thuật."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [1500, 3600, 4260])
    style_table(table, font_size=9.0)

    add_callout(
        doc,
        "SỐ LIỆU HÃNG - CHỈ ĐỂ SÀNG LỌC",
        "Qwen2.5 Technical Report công bố MATH/GSM8K của bản Instruct tăng rõ theo quy mô: 0.5B (34,4/49,6), 1.5B (55,2/73,2), 7B (75,5/91,6). Điều này hợp lý hóa việc thêm Qwen2.5-1.5B vào danh sách ứng viên, nhưng không thay thế test tiếng Việt và test Socratic của nhóm.",
        fill="FFF8E8",
        accent=GOLD,
    )

    doc.add_heading("6. Thiết kế thí nghiệm lựa chọn model", level=1)
    add_body(doc, "Mục tiêu là trả lời câu hỏi: trong cùng giới hạn phần cứng và cùng quy trình fine-tune, model nền nào tạo ra chất lượng nội dung và sư phạm tốt nhất cho từng môn?")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(PIPELINE_IMG), width=Inches(6.35))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run("Hình 1. Pipeline thực nghiệm từ model ứng viên đến Registry và Hybrid Router")
    set_run_font(r, size=9, italic=True, color=GRAY)

    doc.add_heading("6.1. Danh sách ứng viên công bằng", level=2)
    add_table_caption(doc, "Bảng 4. Candidate set đề xuất trong giới hạn capstone")
    table = doc.add_table(rows=1, cols=4)
    for i, t in enumerate(["Ứng viên", "Vai trò", "Môn", "Lý do đưa vào"]):
        table.rows[0].cells[i].text = t
    candidates = [
        ("Qwen/Qwen2.5-0.5B-Instruct", "Baseline hiện tại", "Cả 3", "Nhỏ nhất, đã tích hợp, đo mốc latency/VRAM."),
        ("Qwen/Qwen2.5-1.5B-Instruct", "Scale-up cùng họ", "Cả 3", "Kiểm tra liệu tăng quy mô có sửa lỗi chất lượng đáng kể hay không."),
        ("meta-llama/Llama-3.2-1B-Instruct", "Khác kiến trúc/họ model", "Cả 3", "Đối chứng để tránh chỉ so nội bộ Qwen."),
        ("HuggingFaceTB/SmolLM2-1.7B-Instruct", "Small-model comparator", "Cả 3", "Cùng phân khúc triển khai, open-weight."),
        ("Qwen/Qwen2.5-Math-1.5B-Instruct", "Specialist", "Toán", "Kiểm tra lợi ích chuyên môn hóa so với general model cùng cỡ."),
    ]
    for row in candidates:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [2500, 1700, 950, 4210])
    style_table(table, font_size=8.8, center_cols=[2])
    add_body(doc, "Nếu GPU không đủ, ưu tiên ba model: Qwen2.5-0.5B, Qwen2.5-1.5B và Llama-3.2-1B; chỉ thêm Qwen2.5-Math-1.5B cho Toán. Không cần chạy model 7B/72B vì không cùng điều kiện triển khai.", italic=True)

    doc.add_heading("6.2. Điều kiện phải giữ cố định", level=2)
    for item in [
        "Cùng train/validation/test split; khóa test set trước khi train và không dùng test để chỉnh prompt.",
        "Cùng prompt hệ thống, chat template, max_new_tokens=512, temperature=0,2 và decoding policy.",
        "Cùng LoRA hyperparameters, số epoch, effective batch size, learning rate, seed và early-stopping rule.",
        "Cùng phần cứng, chế độ warm/cold, số lần lặp và thứ tự chạy được randomize.",
        "Ghi model revision/commit, dataset hash, config hash, thời gian chạy và lỗi OOM/Bad Gateway.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("6.3. Ba pha thực nghiệm", level=2)
    num_id = create_numbering_id(doc)
    add_number(doc, "Pha A - Base evaluation: chạy model instruct chưa fine-tune để đo năng lực nền và tránh quy kết mọi khác biệt cho dataset SFT.", num_id)
    add_number(doc, "Pha B - Controlled LoRA/SFT: fine-tune từng model bằng đúng cùng split và hyperparameters. Nếu model cần chat template riêng, ghi đây là khác biệt kiến trúc bắt buộc.", num_id)
    add_number(doc, "Pha C - Locked evaluation: chạy batch test khóa; sau đó mới tích hợp model thắng vào Registry và thực nghiệm end-to-end POOLED/ORACLE/HYBRID.", num_id)

    doc.add_heading("7. Dataset và chỉ số theo từng môn", level=1)
    add_body(doc, "Router Benchmark và Model Evaluation Pack là hai dữ liệu khác mục tiêu. Router Benchmark chấm gold_subject, gold_intent và need_clarification; Model Evaluation Pack chấm chất lượng câu trả lời. Không dùng chung một con số để kết luận cả routing lẫn generation.")
    add_table_caption(doc, "Bảng 5. Evaluation Pack tối thiểu khả thi")
    table = doc.add_table(rows=1, cols=4)
    for i, t in enumerate(["Môn", "Số case khóa", "Nội dung bắt buộc", "Primary metrics"]):
        table.rows[0].cells[i].text = t
    rows = [
        ("Toán", "60", "Tính toán, đại số, lời văn; chuẩn, nhiễu, follow-up; có lời giải và key steps.", "Final-answer accuracy; step correctness; reasoning hallucination."),
        ("Tiếng Anh", "60", "Grammar, vocabulary, reading, explanation; câu hỏi EN và VI→EN; kiểm tra language consistency.", "Task accuracy; language consistency; pedagogical feedback."),
        ("Lịch sử", "60", "Mốc thời gian, nhân vật, nguyên nhân-hệ quả, false premise; có nguồn chuẩn.", "FActScore; unsupported-claim rate; date/entity accuracy; abstention."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [1050, 1100, 4050, 3160])
    style_table(table, font_size=8.8, center_cols=[1])
    add_body(doc, "Khuyến nghị thực tế: 180 cases khóa (60/môn) là mức tối thiểu để so sánh capstone. Tập train 150-200 mẫu/môn có thể dùng để thử nghiệm ban đầu, nhưng phải loại trùng/ngữ nghĩa gần trùng với validation và test. Nếu có thời gian, tăng test lên 100/môn sẽ cho khoảng tin cậy ổn định hơn.", italic=True)

    doc.add_heading("8. Định nghĩa hallucination và điểm Socratic", level=1)
    doc.add_heading("8.1. Hallucination theo môn", level=2)
    add_body(doc, "Không dùng một nhãn hallucination chung cho mọi môn. Mỗi phản hồi được tách thành các claim có thể kiểm chứng, sau đó đối chiếu reference answer hoặc nguồn chuẩn.")
    add_body(doc, "FActScore = số atomic facts được nguồn chuẩn hỗ trợ / tổng số atomic facts có thể kiểm chứng.", bold_lead="FActScore =")
    add_body(doc, "Unsupported-claim rate = số claim không có bằng chứng hoặc mâu thuẫn nguồn / tổng số claim có thể kiểm chứng.", bold_lead="Unsupported-claim rate =")
    add_body(doc, "Với Toán, một lời giải có đáp số đúng nhưng bước trung gian sai vẫn phải gắn reasoning hallucination. Với Lịch sử, sai ngày/thực thể hoặc khẳng định chắc chắn khi đề có false premise đều là lỗi. Với Tiếng Anh, cần tách lỗi kiến thức ngôn ngữ khỏi lỗi đổi ngôn ngữ ngoài yêu cầu.")

    doc.add_heading("8.2. Điểm Socratic hiện tại", level=2)
    add_body(doc, "Công thức thiết kế ban đầu có thể giữ ở RP4: S_Soc = 0,5×A1 + 0,3×A2 + 0,2×A3, trong đó mỗi thành phần nằm trên thang 0-5.")
    add_table_caption(doc, "Bảng 6. Ý nghĩa ba thành phần và căn cứ trọng số")
    table = doc.add_table(rows=1, cols=4)
    for i, t in enumerate(["Mã", "Tiêu chí", "Vai trò", "Căn cứ trọng số ban đầu"]):
        table.rows[0].cells[i].text = t
    rows = [
        ("A1", "Answer Withholding", "Không tiết lộ đáp án trước khi học sinh tự thử.", "0,5 - điều kiện cổng; vi phạm làm mất mục tiêu Socratic."),
        ("A2", "Scaffolding Quality", "Gợi ý từng bước, đúng mức độ.", "0,3 - quyết định chất lượng hỗ trợ sau khi không mớm đáp án."),
        ("A3", "Diagnostic Question", "Đặt câu hỏi kiểm tra hiểu biết/sai lầm.", "0,2 - bổ sung chẩn đoán và duy trì hội thoại."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [800, 2000, 3100, 3460])
    style_table(table, font_size=9.0, center_cols=[0])
    add_callout(doc, "GIỚI HẠN CẦN GHI THẲNG", "Trọng số 0,5/0,3/0,2 là giả thuyết thiết kế, chưa phải trọng số được học từ dữ liệu. RP5 phải chạy sensitivity analysis; nếu thứ hạng model thay đổi mạnh khi đổi trọng số, không được tuyên bố kết quả robust.", fill="FFF3F3", accent=RED)

    doc.add_heading("9. Kiểm định thống kê và quy tắc chọn model", level=1)
    add_body(doc, "Không nên tạo một 'overall score' bằng trọng số tùy ý. Phương án bảo vệ tốt hơn là lựa chọn theo ràng buộc và Pareto.")
    num_id = create_numbering_id(doc)
    add_number(doc, "Định nghĩa trước ngưỡng tối thiểu cho correctness/factuality, Socratic và language consistency. Model không đạt một ngưỡng cổng bị loại dù latency thấp.", num_id)
    add_number(doc, "Trên các model còn lại, tính bootstrap 95% confidence interval. Dùng McNemar cho kết quả nhị phân trên cùng case; Wilcoxon signed-rank cho điểm rubric thứ bậc nếu phân phối không chuẩn.", num_id)
    add_number(doc, "Chỉ tuyên bố model A tốt hơn model B khi chênh lệch có ý nghĩa hoặc đạt tiêu chí non-inferiority định trước. Nếu chất lượng tương đương, chọn model có TTFT, total latency và VRAM thấp hơn.", num_id)
    add_number(doc, "Fine-tuning nên chạy 3 random seeds nếu tài nguyên cho phép. Ít nhất 20% test set cần chấm người độc lập; đo Cohen's kappa giữa người và Gemini Judge, mục tiêu κ ≥ 0,60.", num_id)
    add_body(doc, "Quy tắc cuối: chọn specialist cho một môn chỉ khi specialist cải thiện chất lượng nội dung có ý nghĩa, không làm giảm điểm Socratic dưới ngưỡng và không gây chi phí chuyển model vượt SLA.", bold_lead="Quy tắc cuối:")

    doc.add_heading("10. Tích hợp với Hybrid Router và các report", level=1)
    add_table_caption(doc, "Bảng 7. Vị trí của nội dung trong chuỗi report")
    table = doc.add_table(rows=1, cols=3)
    for i, t in enumerate(["Report", "Nội dung đưa vào", "Không đưa vào"]):
        table.rows[0].cells[i].text = t
    rows = [
        ("RP3 - Literature Review", "Benchmark theo môn/ngôn ngữ; hallucination; khoảng trống không có one-size-fits-all.", "Kết quả test nội bộ chưa chạy."),
        ("RP4 - Methodology", "Candidate set, data split, metrics, công thức, controls, kiểm định và selection rule.", "Bảng điểm thực nghiệm và kết luận model thắng."),
        ("RP5 - Experimental & Results", "Base vs fine-tuned; model theo môn; hallucination; latency/TTFT/VRAM; Router và E2E.", "Tuyên bố chưa được số liệu hỗ trợ."),
        ("RP6 - Discussion/Future Work", "Threats to validity; RAG, curriculum graph, RL router hoặc multi-agent như future work.", "Cam kết RAG/RL đã hiệu quả nếu chưa triển khai."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, [1750, 4200, 3410])
    style_table(table, font_size=8.8)

    doc.add_heading("11. Việc cần làm ngay", level=1)
    tasks = [
        "Chốt candidate set và giới hạn kích thước model với mentor.",
        "Khóa Evaluation Pack 180 cases, có reference answer/key points và nguồn chuẩn cho Lịch sử.",
        "Sửa batch testing để chạy cùng case cho nhiều model, xuất response, latency, TTFT, VRAM và lỗi kỹ thuật.",
        "Chạy Pha A trước; nếu 0.5B thua rõ, không tốn thời gian fine-tune lại model đó nhiều cấu hình.",
        "Chạy controlled LoRA, sau đó locked test và thống kê; lưu đầy đủ config/seed/hash.",
        "Đăng ký model thắng vào Registry và chạy E2E POOLED/ORACLE/HYBRID; đo thêm model-switch latency.",
        "Chuyển toàn bộ bảng kết quả sang RP5; RP4 chỉ giữ protocol, công thức và tiêu chí quyết định.",
    ]
    num_id = create_numbering_id(doc)
    for t in tasks:
        add_number(doc, t, num_id)

    doc.add_heading("12. Các câu hỏi cần mentor chốt", level=1)
    questions = [
        "Giới hạn model được phép triển khai là ≤2B tham số hay có thể dùng model 3B/7B qua external provider?",
        "Có chấp nhận 60 test cases/môn cho vòng capstone đầu tiên không, hay yêu cầu 100 cases/môn?",
        "Tiếng Anh sẽ đánh giá English-only hay có cả yêu cầu giải thích bằng tiếng Việt?",
        "Ai là người chấm human validation subset, và nhóm có thể có hai người chấm độc lập không?",
        "Ngưỡng cổng tối thiểu cho factuality/Socratic/latency là bao nhiêu trước khi xem chi phí?",
        "Có chấp nhận Qwen2.5-Math-1.5B làm specialist riêng cho Toán trong khi hai môn còn lại dùng general model không?",
    ]
    for q in questions:
        add_bullet(doc, q)

    doc.add_heading("13. Kịch bản trình bày mentor trong 8-10 phút", level=1)
    script = [
        ("Phút 0-1", "Nêu vấn đề: model hiện tại chọn theo khả năng triển khai, chưa được chứng minh tốt nhất theo môn."),
        ("Phút 1-3", "Trình bày 4 nguồn cốt lõi: VMLU, K-12EduBench/EduBench, MathTutorBench và FActScore."),
        ("Phút 3-5", "Kết luận literature: không có one-size-fits-all; correctness, pedagogy và hallucination phải đo tách."),
        ("Phút 5-7", "Trình bày pipeline Hình 1 và candidate set trong cùng phân khúc tài nguyên."),
        ("Phút 7-8", "Trình bày Evaluation Pack 180 cases và quy tắc Pareto/non-inferiority."),
        ("Phút 8-10", "Xin mentor chốt giới hạn model, số case, ngôn ngữ đánh giá và người chấm."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Thời gian"
    table.rows[0].cells[1].text = "Nội dung nói"
    for row in script:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = row
    set_table_geometry(table, [1450, 7910])
    style_table(table, font_size=9.2, center_cols=[0])

    doc.add_page_break()
    doc.add_heading("Tài liệu tham khảo", level=1)
    add_body(doc, "Nhóm A - Bài đã phản biện (ưu tiên trích dẫn trong RP3/RP4)", bold_lead="Nhóm A - Bài đã phản biện")
    refs = [
        ("[1]", "Bui, C. T., et al. (2025). VMLU Benchmarks: A comprehensive benchmark toolkit for Vietnamese LLMs. ACL 2025, 11495-11515. DOI: 10.18653/v1/2025.acl-long.563.", "https://aclanthology.org/2025.acl-long.563/"),
        ("[2]", "Ye, Y., et al. (2026). K-12EduBench: A Benchmark for Evaluating Large Language Models' Knowledge, Problem-Solving, and Educational Goal Cognition in K-12 Education. AAAI 2026, 40(40), 34459-34466. DOI: 10.1609/aaai.v40i40.40744.", "https://ojs.aaai.org/index.php/AAAI/article/view/40744"),
        ("[3]", "Xu, B., et al. (2026). EduBench: A Comprehensive Benchmarking Dataset for Evaluating Large Language Models in Diverse Educational Scenarios. ACL 2026, 21615-21645. DOI: 10.18653/v1/2026.acl-long.987.", "https://aclanthology.org/2026.acl-long.987/"),
        ("[4]", "Wang, Y., et al. (2024). MMLU-Pro: A More Robust and Challenging Multi-Task Language Understanding Benchmark. NeurIPS 2024, Datasets and Benchmarks Track. DOI: 10.52202/079017-3018.", "https://proceedings.neurips.cc/paper_files/paper/2024/hash/ad236edc564f3e3156e1b2feafb99a24-Abstract-Datasets_and_Benchmarks_Track.html"),
        ("[5]", "Hou, X., et al. (2024). E-EVAL: A Comprehensive Chinese K-12 Education Evaluation Benchmark for Large Language Models. Findings of ACL 2024. DOI: 10.18653/v1/2024.findings-acl.462.", "https://aclanthology.org/2024.findings-acl.462/"),
        ("[6]", "Macina, J., et al. (2025). MathTutorBench: A Benchmark for Measuring Open-ended Pedagogical Capabilities of LLM Tutors. EMNLP 2025, 204-221. DOI: 10.18653/v1/2025.emnlp-main.11.", "https://aclanthology.org/2025.emnlp-main.11/"),
        ("[7]", "Schorcht, S., Müller, F. A., & Buchholtz, N. (2026). No one-size-fits-all: a study of prompt techniques and large language models to enhance AI's mathematics educational quality. ZDM - Mathematics Education. DOI: 10.1007/s11858-026-01784-6.", "https://link.springer.com/article/10.1007/s11858-026-01784-6"),
        ("[8]", "Min, S., et al. (2023). FActScore: Fine-grained Atomic Evaluation of Factual Precision in Long Form Text Generation. EMNLP 2023, 12076-12100. DOI: 10.18653/v1/2023.emnlp-main.741.", "https://aclanthology.org/2023.emnlp-main.741/"),
        ("[9]", "Li, J., et al. (2023). HaluEval: A Large-Scale Hallucination Evaluation Benchmark for Large Language Models. EMNLP 2023. DOI: 10.18653/v1/2023.emnlp-main.397.", "https://aclanthology.org/2023.emnlp-main.397/"),
        ("[10]", "Lin, S., Hilton, J., & Evans, O. (2022). TruthfulQA: Measuring How Models Mimic Human Falsehoods. ACL 2022. DOI: 10.18653/v1/2022.acl-long.229.", "https://aclanthology.org/2022.acl-long.229/"),
        ("[11]", "Wang, X., et al. (2024). SciBench: Evaluating College-Level Scientific Problem-Solving Abilities of Large Language Models. ICML 2024, PMLR 235:50622-50649.", "https://proceedings.mlr.press/v235/wang24z.html"),
        ("[12]", "Ji, Z., et al. (2023). Survey of Hallucination in Natural Language Generation. ACM Computing Surveys, 55(12). DOI: 10.1145/3571730.", "https://doi.org/10.1145/3571730"),
    ]
    for num, text, url in refs:
        p = doc.add_paragraph(style="Reference")
        r = p.add_run(num + " ")
        set_run_font(r, size=9.5, bold=True, color=NAVY)
        r = p.add_run(text + " ")
        set_run_font(r, size=9.5)
        add_hyperlink(p, "Nguồn", url)

    add_body(doc, "Nhóm B - Technical report dùng để sàng lọc ứng viên, không phải bằng chứng độc lập", bold_lead="Nhóm B - Technical report")
    tech_refs = [
        ("[13]", "Yang, A., et al. (2024). Qwen2.5 Technical Report. arXiv:2412.15115.", "https://arxiv.org/abs/2412.15115"),
        ("[14]", "Yang, A., et al. (2024). Qwen2.5-Math Technical Report: Toward Mathematical Expert Model via Self-Improvement. arXiv:2409.12122.", "https://arxiv.org/abs/2409.12122"),
    ]
    for num, text, url in tech_refs:
        p = doc.add_paragraph(style="Reference")
        r = p.add_run(num + " ")
        set_run_font(r, size=9.5, bold=True, color=NAVY)
        r = p.add_run(text + " ")
        set_run_font(r, size=9.5)
        add_hyperlink(p, "Nguồn", url)

    add_body(doc, "Ghi chú học thuật: Các nguồn arXiv mới như OpenLearnLM, K12-KGraph và Theorem-SFT chỉ nên đặt ở Future Work cho đến khi có phản biện hoặc tái lập độc lập. Các trang tổng hợp benchmark, blog và press release đã bị loại khỏi phần bằng chứng chính.", italic=True)

    # Document metadata
    props = doc.core_properties
    props.title = "Báo cáo cơ sở khoa học lựa chọn pretrained model theo môn học"
    props.subject = "Hybrid Multi-SLM Router - báo cáo mentor"
    props.author = "SEP490 Group 36"
    props.keywords = "LLM, model selection, education, hallucination, hybrid router, SEP490"
    props.comments = "Generated as a mentor research brief; results must be validated experimentally."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    build_document()
