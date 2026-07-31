from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\Sep_G36\SEP490_G36\SEP490_WEBAPP-main\docs\Ke_hoach_trien_khai_va_test_nghien_cuu_Hybrid_Router.docx")
BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
GRAY = "F2F4F7"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    ind = tbl_pr.find(qn('w:tblInd'))
    if ind is None:
        ind = OxmlElement('w:tblInd')
        tbl_pr.append(ind)
    ind.set(qn('w:w'), '120')
    ind.set(qn('w:type'), 'dxa')
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(widths[index]))
            tc_w.set(qn('w:type'), 'dxa')


def set_font(run, name='Calibri', size=11, color=None, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ('Heading 1', 16, BLUE, 18, 10),
        ('Heading 2', 13, BLUE, 14, 7),
        ('Heading 3', 12, DARK, 10, 5),
    ]:
        style = styles[name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for list_name in ['List Bullet', 'List Number']:
        style = styles[list_name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('KẾ HOẠCH TRIỂN KHAI VÀ TEST NGHIÊN CỨU')
    set_font(r, size=20, color=DARK, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(16)
    r = p2.add_run('Hybrid Multi-Model Router cho AI Tutor Socratic')
    set_font(r, size=14, color=BLUE, bold=True)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run('Bản tiếng Việt dùng để triển khai, ghi số liệu và hoàn thiện RP4–RP6')
    set_font(r, size=10, color='555555', italic=True)


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items, numbered=False):
    for item in items:
        p = doc.add_paragraph(style='List Number' if numbered else 'List Bullet')
        r = p.add_run(item)
        set_font(r)


def add_note(doc, label, text, color=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(label + ': ')
    set_font(r, bold=True, color=DARK if color != 'FDECEC' else RED)
    r = p.add_run(text)
    set_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        r = p.add_run(header)
        set_font(r, size=10, color=DARK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            r = p.add_run(str(value))
            set_font(r, size=9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def main():
    doc = Document()
    style_document(doc)
    add_title(doc)
    add_note(doc, 'Mục tiêu của tài liệu', 'Dùng như checklist điều hành toàn bộ phần còn lại của nghiên cứu: chuẩn bị dữ liệu, train, đăng ký model, benchmark Router, Model Eval/RP5, đo latency và đưa kết quả vào báo cáo. Không dùng tài liệu này để thay thế nội dung học thuật của RP1–RP7.', LIGHT)

    h(doc, '1. Phạm vi và kiến trúc phải giữ cố định')
    add_para(doc, 'V1 là baseline một mô hình tutor dùng chung cho mọi câu hỏi. V2 là hệ thống Hybrid Multi-Model: Rule Router xử lý case rõ, LLM-as-Router xử lý case khó, sau đó chọn đúng một response model để sinh câu trả lời Socratic.')
    add_table(doc, ['Thành phần', 'Vai trò', 'Không được nhầm với'], [
        ('GENERAL', 'Mô hình pooled dùng cho câu chung, xã giao hoặc baseline B.', 'Không phải Router.'),
        ('Specialist', 'ENGLISH, MATH, HISTORY; mỗi model được train riêng theo môn.', 'Không phải Gemini Judge.'),
        ('Hybrid Router', 'Chọn subject, intent, clarification và model.', 'Không sinh câu trả lời cuối.'),
        ('Oracle', 'Dùng gold_subject để gọi specialist đúng; upper bound về model selection.', 'Không phải Router thực tế.'),
        ('Gemini Judge', 'Chấm post-hoc chất lượng phản hồi trong RP5.', 'Không tham gia chọn model khi inference.'),
    ], [1900, 4800, 2660])
    add_note(doc, 'Điều kiện scope', 'Không mở rộng sang RL/RCT trong capstone hiện tại. Tập trung vào routing, LoRA specialists, chất lượng Socratic và trade-off latency/cost.', 'F2F4F7')

    h(doc, '2. Dataset cần chốt trước khi train')
    add_para(doc, 'Bộ mới đã được tạo riêng theo ba môn. Dataset cũ Physics/Chemistry vẫn còn để tham chiếu, nhưng không trộn vào thí nghiệm chính mới.')
    add_table(doc, ['File', 'Số mẫu', 'Mục đích', 'Ghi chú'], [
        ('specialist_english_train_v2.json', '200', 'Train specialist ENGLISH', 'messages + system/user/assistant'),
        ('specialist_math_train_v2.json', '200', 'Train specialist MATH', 'messages + system/user/assistant'),
        ('specialist_history_train_v2.json', '200', 'Train specialist HISTORY', 'messages + system/user/assistant'),
        ('general_tutor_train_v1.json', '600', 'GENERAL / pooled baseline', 'Giữ riêng, không gán gold subject chuyên biệt'),
        ('router_train_v2.json', '90', 'Calibration/training Router', '30 mẫu mỗi môn'),
        ('router_test_v2.json', '60', 'Hold-out Router Benchmark', '20 mẫu mỗi môn, không tune threshold'),
    ], [3150, 850, 2600, 2760])
    h(doc, '2.1. Kiểm tra dataset trước upload', 2)
    add_bullets(doc, [
        'Đếm đúng 200 mẫu trong từng specialist train và 60 case trong router_test_v2.',
        'Kiểm tra ba nhãn chính: ENGLISH, MATH, HISTORY; không có nhãn viết sai hoặc khoảng trắng thừa.',
        'Mỗi router case phải có gold_subject, gold_intent và gold_need_clarification.',
        'Giữ challenge_type: standard, typo, abbreviation, follow_up, ambiguous.',
        'Không dùng router_test_v2 để chọn learning rate, epochs hoặc threshold.',
        'Lưu manifest, checksum hoặc tên version dataset để tái lập thí nghiệm.',
    ])
    add_note(doc, 'Cảnh báo dữ liệu', 'Các file hiện tại là synthetic development data. Có thể dùng để kiểm tra pipeline, nhưng trong RP phải ghi rõ giới hạn và không gọi đây là dữ liệu hội thoại học sinh thật.', 'FDECEC')

    h(doc, '3. Train specialist models')
    add_para(doc, 'Train GENERAL và ba specialist bằng cùng base model, cùng prompt format và cùng quy tắc lưu log. Chỉ thay dataset/subject; không thay đồng thời nhiều biến trong thí nghiệm chính.')
    h(doc, '3.1. Cấu hình khởi đầu', 2)
    add_table(doc, ['Hyperparameter', 'Giá trị', 'Lý do cần ghi trong RP4'], [
        ('Epochs', '3', 'Dataset nhỏ; giảm nguy cơ overfitting.'),
        ('Learning rate', '5e-5', 'Mức trung bình cho LoRA fine-tuning; cần xác nhận bằng pilot.'),
        ('LoRA rank r', '16', 'Đủ capacity cho đặc trưng môn học nhưng không quá nặng VRAM.'),
        ('LoRA alpha', '32', 'Giữ scale cập nhật adapter ở mức ổn định.'),
        ('LoRA dropout', '0.05', 'Regularization nhẹ cho dữ liệu tổng hợp.'),
        ('Batch / accumulation', '1 / 4', 'Phù hợp giới hạn GPU Colab.'),
        ('Max length', '1024', 'Đủ chứa system prompt, history ngắn và câu trả lời.'),
        ('Seed', '3407', 'Tạo run chính tái lập; nên chạy thêm seed 42/1234 nếu có thời gian.'),
    ], [2300, 1600, 5460])
    h(doc, '3.2. Pilot và sensitivity test', 2)
    add_para(doc, 'Không cần grid search lớn. Cần một thí nghiệm nhỏ có kiểm soát để chứng minh tham số không được điền tùy ý.')
    add_table(doc, ['Cấu hình', 'Epochs', 'Learning rate', 'LoRA r', 'Dùng để'], [
        ('Conservative', '2', '2e-5', '8', 'Kiểm tra underfitting / ổn định.'),
        ('Proposed', '3', '5e-5', '16', 'Cấu hình chính dự kiến.'),
        ('Capacity-high', '5', '1e-4', '32', 'Kiểm tra thêm capacity và overfitting.'),
    ], [2100, 1250, 1800, 1100, 3110])
    add_bullets(doc, [
        'Giữ cố định base model, dataset split, prompt, seed và max length khi so sánh.',
        'Chọn cấu hình bằng validation loss và chất lượng validation; không chọn bằng router_test_v2.',
        'Lưu train loss, eval loss, thời gian train, peak VRAM, checkpoint và config JSON.',
        'Nếu kết quả giữa các seed dao động lớn, báo cáo standard deviation thay vì chỉ báo cáo một run.',
    ])
    add_note(doc, 'Tiêu chí chọn', 'Cấu hình cuối không nhất thiết là cấu hình có loss thấp nhất; phải cân bằng factuality, Socratic quality, latency và khả năng tái lập.', LIGHT)

    h(doc, '4. Hugging Face và Model Registry')
    add_bullets(doc, [
        'Push từng model vào một repository riêng; tên repository phải xác định được subject.',
        'Trong Registry, đăng ký đúng hfRepoId/hf_hub_id, không chỉ ghi tên hiển thị.',
        'Gán subject chính xác: English → ENGLISH; Math → MATH; History → HISTORY; General → GENERAL.',
        'Chọn đúng version Active/USE trước khi chạy Router Benchmark.',
        'Nếu còn registry Physics/Chemistry cũ, tắt routerEnabled hoặc không đưa vào model map của thí nghiệm mới.',
        'Lưu lại HF repository, commit/version, dataset, training config và training loss.',
    ])
    add_table(doc, ['Kiểm tra', 'Đạt khi'], [
        ('HF repository', 'Mở được repository và có config/tokenizer/adapter files.'),
        ('hfRepoId', 'Không rỗng; dạng owner/repository.'),
        ('Registry subject', 'Khớp nhãn dataset và model map.'),
        ('Active version', 'Có trạng thái USE/Active trước khi test.'),
        ('GPU load', 'Load thành công và không báo Missing hf_model_id/Model mismatch.'),
    ], [2600, 6760])

    h(doc, '5. Router Benchmark – RP4')
    add_para(doc, 'Upload router_test_v2.json. Đây là thí nghiệm đo năng lực chọn tuyến, không chấm chất lượng câu trả lời cuối.')
    h(doc, '5.1. Các mode phải chạy', 2)
    add_table(doc, ['Mode', 'Ý nghĩa', 'Số lần lặp'], [
        ('RULE', 'Nhanh, không gọi LLM; chỉ dùng lexical/rule signal.', '3'),
        ('LLM', 'LLM-as-Router cho mọi case; làm baseline accuracy/latency.', '3'),
        ('HYBRID', 'Rule trước, LLM fallback khi low-confidence/ambiguous.', '3'),
        ('Ablations', 'NO_HISTORY, NO_PREVIOUS_SUBJECT, STRICT, LENIENT.', '3'),
    ], [1700, 5900, 1760])
    h(doc, '5.2. Chỉ số phải lưu', 2)
    add_bullets(doc, [
        'Subject Accuracy = số subject đúng / tổng số case.',
        'Macro-F1 = trung bình F1 của ENGLISH, MATH và HISTORY.',
        'Intent Accuracy và Exact Match của subject + intent + clarification.',
        'Wrong-route count, clarification accuracy và stability giữa các repeats.',
        'Router latency, LLM call rate, router tokens và chi phí router.',
        'Per-subject metrics và danh sách lỗi theo challenge_type.',
    ])
    add_note(doc, 'Không được suy ra từ RP4', 'Điểm Socratic, factuality, overall và latency sinh câu trả lời không được suy ra từ Router Benchmark. Phải chạy cùng replay qua Model Eval/RP5.', 'FDECEC')

    h(doc, '6. Model Eval / RP5 – chất lượng end-to-end')
    add_para(doc, 'Dùng cùng một bộ prompt/replay cho ba điều kiện để tách năng lực model khỏi lỗi Router.')
    add_table(doc, ['Điều kiện', 'Cách chạy', 'Câu hỏi nghiên cứu'], [
        ('B. POOLED', 'GENERAL nhận mọi prompt.', 'Một model dùng chung đạt chất lượng thế nào?'),
        ('C. ORACLE', 'Gold subject gọi specialist đúng.', 'Trần trên của việc chọn specialist là bao nhiêu?'),
        ('D. HYBRID', 'Router thực tế chọn model.', 'Hybrid có tiệm cận Oracle với chi phí thấp hơn không?'),
    ], [1900, 4200, 3260])
    h(doc, '6.1. Judge rubric', 2)
    add_para(doc, 'Judge chấm từng response theo thang 0–5 hoặc 1–5 đã khóa trước khi chạy. Các nhóm điểm cần giữ nhất quán: A – Socratic pedagogy, B – factuality/grade appropriateness, C – coherence/tone, D – hallucination/robustness.')
    add_para(doc, 'Công thức tổng quát cần ghi rõ là:')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Overall = wA·A + wB·B + wC·C + wD·D,   với   wA + wB + wC + wD = 1')
    set_font(r, size=11, color=DARK, bold=True)
    add_bullets(doc, [
        'Trọng số ban đầu là giả thuyết engineering, không được gọi là đã được xác nhận nếu chưa có calibration.',
        'Ưu tiên A cao vì đề tài đánh giá tutor Socratic; B vẫn phải có hard constraint để câu trả lời sai không đạt điểm cao.',
        'Dùng 30–50 mẫu calibration có ít nhất hai người chấm để so sánh với Overall của Judge.',
        'Chọn trọng số bằng expert elicitation/AHP và Spearman correlation với human score; giữ test set độc lập.',
        'Báo cáo sensitivity profiles: pedagogy-first, balanced và factuality-first.',
    ])
    add_note(doc, 'Lỗi cần kiểm tra trong code', 'Nếu rubric chỉ có D1 nhưng công thức còn dùng D2, phải sửa về D = D1 hoặc bổ sung tiêu chí D2 trước khi chạy RP5 chính thức.', 'FDECEC')

    h(doc, '7. Test tốc độ và chuyển model')
    add_para(doc, 'Để khẳng định tốc độ ổn định, phải đo end-to-end chứ không chỉ lấy Router latency.')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('T_total = T_router + T_model-switch + T_generation')
    set_font(r, size=12, color=DARK, bold=True)
    add_table(doc, ['Kịch bản', 'Chuỗi model', 'Mục tiêu'], [
        ('Warm same-model', 'MATH → MATH → MATH', 'Độ trễ sinh bình thường khi model đã nóng.'),
        ('Two-model switch', 'GENERAL → MATH → GENERAL', 'Kiểm tra chuyển tuyến phổ biến.'),
        ('Four-model sequence', 'GENERAL → MATH → HISTORY → ENGLISH', 'Kiểm tra vượt giới hạn slot/cache.'),
        ('Return after eviction', 'MATH → HISTORY → ENGLISH → MATH', 'Đo cold reload khi model bị thay.'),
        ('Cold start', 'Model chưa load → inference', 'Đo thời gian tải lần đầu.'),
        ('Follow-up', 'MATH → follow-up mơ hồ', 'Kiểm tra history/previous_subject và latency.'),
    ], [2200, 3000, 4160])
    add_bullets(doc, [
        'Chạy mỗi kịch bản tối thiểu 10 lần; tốt hơn là 30 lần nếu GPU ổn định.',
        'Lưu mean, median/P50, P95, standard deviation, TTFT và total latency.',
        'Lưu model trước/sau, cache hit, switch_required, switch latency, số lần reload và lỗi Model mismatch.',
        'Tách cold latency và warm latency; không trộn hai loại vào một trung bình.',
        'Tính CV = standard deviation / mean để đánh giá độ ổn định.',
    ])
    add_note(doc, 'Rủi ro kiến trúc hiện tại', 'GPU service hiện có 2 slot và mỗi slot giữ một model. Benchmark có slot acquisition, nhưng chat trực tiếp phải được kiểm tra model switching riêng; nếu slot chứa model khác, có thể phát sinh Model mismatch hoặc cold reload.', 'FDECEC')

    h(doc, '8. Phân tích thống kê và cách kết luận')
    add_bullets(doc, [
        'Không kết luận từ một case hoặc một run duy nhất.',
        'Với Router, báo cáo accuracy/F1 theo subject và challenge_type; không chỉ báo cáo pooled average.',
        'Với RP5, báo cáo mean ± standard deviation cho A/B/C/D và overall của B/C/D.',
        'Với latency, báo cáo P50/P95 để phản ánh trải nghiệm người dùng và tail latency.',
        'Nếu có đủ repeats, dùng paired comparison trên cùng prompt giữa B, C và D; ghi rõ test thống kê đã dùng.',
        'Nếu chưa đủ mẫu để kiểm định ý nghĩa, dùng ngôn ngữ thận trọng: observed difference, not statistically conclusive.',
    ])
    add_table(doc, ['Kết quả quan sát', 'Diễn giải hợp lý'], [
        ('Hybrid gần Oracle', 'Router chọn specialist khá tốt; chất lượng giới hạn chủ yếu nằm ở model/data.'),
        ('Hybrid thấp hơn Oracle rõ rệt', 'Có lỗi định tuyến hoặc lỗi history/clarification.'),
        ('Oracle cũng thấp', 'Specialist hoặc dữ liệu train chưa tốt; không thể đổ lỗi cho Router.'),
        ('RULE nhanh nhưng F1 thấp', 'Trade-off speed/accuracy; cần dùng Hybrid fallback.'),
        ('LLM accuracy cao nhưng latency/cost cao', 'Có cơ sở định lượng cho thiết kế cascade.'),
    ], [3000, 6360])

    h(doc, '9. Đưa kết quả vào các Report')
    add_table(doc, ['Report', 'Nội dung cần đưa'], [
        ('RP1 – Introduction', 'V1 one-model problem, motivation, mục tiêu V2 và scope capstone.'),
        ('RP2 – Proposal and Plan', 'Research questions, work packages, dataset/model plan, timeline và rủi ro.'),
        ('RP3 – Literature Review', 'Related routing, LoRA/parameter-efficient tuning, LLM-as-router, Socratic tutoring; gap cụ thể.'),
        ('RP4 – Methodology', 'Pipeline V1/V2, dataset schema, train config rationale, pilot/sensitivity, Router metrics, RP5 design, latency formula và test protocol.'),
        ('RP5 – Experimental and Results', 'Bảng train pilot, Router RULE/LLM/HYBRID/ablations, B/C/D Model Eval, latency P50/P95, error analysis.'),
        ('RP6 – Discussion and Future Work', 'Ý nghĩa kết quả, validity threats, synthetic data limitation, slot/cold-start, LoRA adapter switching, RL/RCT future work.'),
        ('RP7 – Final Report', 'Tổng hợp RQ → evidence → conclusion; không đưa claim vượt quá số liệu.'),
    ], [1900, 7460])

    h(doc, '10. Checklist thực thi theo thứ tự')
    add_bullets(doc, [
        '☐ Chốt scope ba specialist ENGLISH/MATH/HISTORY và GENERAL pooled baseline.',
        '☐ Kiểm tra dataset counts, labels, schema và hold-out separation.',
        '☐ Chạy pilot ba cấu hình train; lưu config, loss, latency train và checkpoint.',
        '☐ Chọn cấu hình bằng validation; train final models với seed đã ghi.',
        '☐ Push HF và xác nhận hfRepoId không rỗng.',
        '☐ Register model với subject/Active version chính xác.',
        '☐ Upload router_test_v2; chạy RULE, LLM, HYBRID và ablations, mỗi mode 3 repeats.',
        '☐ Export Router results và lỗi theo case/subject/challenge.',
        '☐ Chạy POOLED, ORACLE, HYBRID trên cùng replay qua Model Eval/Gemini Judge.',
        '☐ Chạy latency scenarios warm/cold/switch; lưu raw logs.',
        '☐ Tính mean, P50, P95, std, CV và paired comparisons nếu đủ repeats.',
        '☐ Sửa rubric D1/D2 và chốt trọng số sau calibration trước khi chạy kết quả chính.',
        '☐ Cập nhật RP4 methodology trước, sau đó điền số liệu vào RP5 và thảo luận ở RP6.',
    ])

    h(doc, '11. Bộ hồ sơ phải lưu để tái lập')
    add_table(doc, ['Nhóm', 'File/thông tin cần lưu'], [
        ('Dataset', 'Các JSON v2, manifest, checksum, ngày tạo, ghi chú synthetic.'),
        ('Training', 'Training config, seed, dataset split, train/eval loss, checkpoint, thời gian và VRAM.'),
        ('HF/Registry', 'Repository ID, commit/version, subject, Active version, model map.'),
        ('RP4', 'Benchmark input, mode, repeats, raw decisions, metrics, error cases.'),
        ('RP5', 'Replay input, B/C/D outputs, Judge raw scores, rubric version, weights.'),
        ('Latency', 'Raw timestamps, model state, cache hit, switch/reload, TTFT, tokens, errors.'),
        ('Reports', 'Bảng số liệu cuối, script phân tích, phiên bản code và ngày chạy.'),
    ], [1900, 7460])
    add_note(doc, 'Điều kiện được phép kết luận', 'Chỉ kết luận Hybrid tốt hơn hoặc nhanh hơn khi cùng dataset/replay, cùng prompt/evaluation protocol, model version đã cố định và có log đủ để kiểm tra lại.', LIGHT)

    h(doc, '12. Trạng thái hiện tại của code/dataset')
    add_para(doc, 'Đã có: dataset v2 ba môn, router train/test v2, mapping Router cho ENGLISH/HISTORY, backend build thành công và frontend/backend có thể khởi động. Còn phải làm: train/push/register ba model mới, chạy đầy đủ RP4/RP5, đo latency raw và điền số liệu vào reports.')
    add_para(doc, 'Đường dẫn dataset chính: D:\\Sep_G36\\SEP490_G36\\SEP490_WEBAPP-main\\docs\\datasets\\. Tài liệu này không ghi kết quả giả; các ô số liệu RP5 và latency chỉ được điền sau khi chạy thực nghiệm.')

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('SEP490 – Research Execution Plan')
    set_font(r, size=9, color='777777')
    doc.core_properties.title = 'Kế hoạch triển khai và test nghiên cứu Hybrid Router'
    doc.core_properties.subject = 'Dataset, training, Router Benchmark, Model Eval và latency testing'
    doc.core_properties.author = 'SEP490 Research Team'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
